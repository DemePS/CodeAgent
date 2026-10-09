"""How a workbook looks: view_excel (see a sheet or range as a picture) and format_excel (bold header,
widths, wrap, fill, borders, freeze panes, filters), with either backend (see documents.excel_backend).

- xlwings: Excel renders the sheet (the PDF it would print, charts and pictures included) or a
  range (a PNG), and applies the formatting itself.
- openpyxl: the sheet is drawn as a web page (column widths, fonts, fills, borders, wrap, merged
  cells) and photographed with the screenshot_page browser; charts and pictures are not drawn (they
  are listed). Formatting is written with openpyxl, which rewrites the file (the person is warned
  when that loses charts or pictures, as with edit_excel).
"""

from __future__ import annotations

import html
import io
import re
import tempfile
from pathlib import Path

from .. import backups, state
from ..common import ToolError, display, is_protected, rel_name
from ..config import EXCEL_VIEW_MAX_CELLS, MAX_IMAGE_BYTES, PDF_MAX_VISUAL_PAGES, uses_deepseek
from .documents import (
    excel_backend,
    excel_path,
    image_block,
    load_workbook,
    lossy_features,
    render_pdf_pages,
    save_openpyxl,
    sheet_objects,
    show_cell,
)

COLOR = re.compile(r"^#?[0-9A-Fa-f]{6}$")
CELL = re.compile(r"^[A-Z]{1,3}[1-9][0-9]{0,6}$")
RANGE = re.compile(r"^[A-Z]{1,3}[1-9][0-9]{0,6}(:[A-Z]{1,3}[1-9][0-9]{0,6})?$|^[A-Z]{1,3}:[A-Z]{1,3}$|^[1-9][0-9]*:[1-9][0-9]*$")


# --- view_excel ----------------------------------------------------------------------------------------

def tool_view_excel(path: str, sheet: str | None = None, range: str | None = None) -> list:
    from openpyxl.utils import range_boundaries

    p = excel_path(path, readable=True)
    wb = load_workbook(p)
    name = sheet or wb.sheetnames[0]
    if name not in wb.sheetnames:
        raise ToolError(f"No sheet {name!r}. Sheets: {', '.join(wb.sheetnames)}")
    address = range.upper().replace("$", "").strip() if range else None
    if address:
        try:
            range_boundaries(address)
        except ValueError:
            raise ToolError(f"Invalid range {range!r}; use e.g. 'A1:F40'.")
    if excel_backend() == "xlwings":
        from . import excel_xl
        data, media_type = excel_xl.render(p, name, address)
        state.ui.status(f"[excel] {rel_name(p)} {name}{' ' + address if address else ''} rendered by Excel")
        label = f"{display(p)} -- sheet {name}{' range ' + address if address else ''}, as Excel shows it"
        if media_type == "application/pdf":
            return [{"type": "text", "text": label + " (the pages Excel would print; charts and pictures included)"}] \
                + pdf_block(data, display(p))
        return [{"type": "text", "text": label}, checked_image(data, "image/png")]
    ws = wb[name]
    others = sheet_objects(ws)
    results = load_workbook(p, data_only=True)[name]  # a formula shows its result, as in Excel
    data = screenshot(sheet_html(ws, address, results))
    state.ui.status(f"[excel] {rel_name(p)} {name}{' ' + address if address else ''} drawn (without Excel)")
    note = (f"{display(p)} -- sheet {name}{' range ' + address if address else ''}, drawn without Excel: cells, "
            "widths, fonts, fills, borders and merged cells as saved"
            + ("; not drawn: " + "; ".join(others) if others else "") + ".")
    return [{"type": "text", "text": note}, checked_image(data, "image/png")]


def checked_image(data: bytes, media_type: str) -> dict:
    if len(data) > MAX_IMAGE_BYTES:
        raise ToolError(f"The picture is {len(data):,} bytes (limit {MAX_IMAGE_BYTES:,}): view a smaller range.")
    return image_block(data, media_type)


def pdf_block(data: bytes, title: str) -> list:
    import base64

    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(io.BytesIO(data))
    count = len(reader.pages)
    if uses_deepseek():  # images only: render the first pages to pictures
        shown = min(count, PDF_MAX_VISUAL_PAGES)
        with tempfile.TemporaryDirectory() as folder:
            pdf = Path(folder) / "sheet.pdf"
            pdf.write_bytes(data)
            pages = render_pdf_pages(pdf, list(range(1, shown + 1)))
        return [{"type": "text", "text": f"Each page below is a picture of the printed sheet ({shown} of {count} page(s))."}] \
            + [checked_image(png, media_type) for png, media_type in pages]
    if count > PDF_MAX_VISUAL_PAGES:  # the first pages only
        writer = PdfWriter()
        for page in reader.pages[:PDF_MAX_VISUAL_PAGES]:
            writer.add_page(page)
        buffer = io.BytesIO()
        writer.write(buffer)
        data = buffer.getvalue()
    return [{"type": "document", "title": title, "context": f"{min(count, PDF_MAX_VISUAL_PAGES)} of {count} page(s)",
             "source": {"type": "base64", "media_type": "application/pdf",
                        "data": base64.b64encode(data).decode("ascii")}}]


def argb(color) -> str | None:
    """'#RRGGBB' from an openpyxl color (theme and indexed colors are not resolved: None)."""
    value = getattr(color, "rgb", None)
    if isinstance(value, str) and len(value) in (6, 8) and getattr(color, "type", "rgb") == "rgb":
        return "#" + value[-6:]
    return None


def sheet_html(ws, address: str | None, results=None) -> str:
    """The sheet (or range) as an HTML table, styled from the cells: what the person sees, roughly."""
    from openpyxl.utils import get_column_letter, range_boundaries

    if address:
        min_col, min_row, max_col, max_row = range_boundaries(address)
        min_col, min_row = min_col or 1, min_row or 1
        max_col, max_row = max_col or ws.max_column, max_row or ws.max_row
    else:
        min_col, min_row, max_col, max_row = 1, 1, max(ws.max_column, 1), max(ws.max_row, 1)
    max_col = min(max_col, min_col + 39)
    max_row = min(max_row, min_row + max(1, EXCEL_VIEW_MAX_CELLS // max(1, max_col - min_col + 1)) - 1)
    spans, hidden = {}, set()
    for merged in ws.merged_cells.ranges:
        spans[(merged.min_row, merged.min_col)] = (merged.max_row - merged.min_row + 1, merged.max_col - merged.min_col + 1)
        hidden |= {(r, c) for r in range(merged.min_row, merged.max_row + 1)
                   for c in range(merged.min_col, merged.max_col + 1)} - {(merged.min_row, merged.min_col)}
    cols = []
    for c in range(min_col, max_col + 1):
        width = ws.column_dimensions[get_column_letter(c)].width or 8.43
        cols.append(f'<col style="width:{round(width * 7 + 5)}px">')
    rows = []
    for r in range(min_row, max_row + 1):
        height = ws.row_dimensions[r].height
        cells = [f'<th class="rh">{r}</th>']
        for c in range(min_col, max_col + 1):
            if (r, c) in hidden:
                continue
            cell = ws.cell(row=r, column=c)
            style = []
            font = cell.font
            if font is not None:
                if font.b:
                    style.append("font-weight:bold")
                if font.i:
                    style.append("font-style:italic")
                if font.sz:
                    style.append(f"font-size:{float(font.sz) * 1.33:.0f}px")
                if (color := argb(font.color)) and color != "#000000":
                    style.append(f"color:{color}")
            if cell.fill is not None and cell.fill.fill_type == "solid" and (color := argb(cell.fill.fgColor)):
                style.append(f"background:{color}")
            align = cell.alignment
            if align is not None:
                if align.horizontal in ("left", "center", "right"):
                    style.append(f"text-align:{align.horizontal}")
                if align.wrap_text:
                    style.append("white-space:pre-wrap")
                if align.vertical in ("top", "center", "bottom"):
                    style.append(f"vertical-align:{'middle' if align.vertical == 'center' else align.vertical}")
            border = cell.border
            for side in ("top", "right", "bottom", "left"):
                edge = getattr(border, side, None) if border is not None else None
                if edge is not None and edge.style:
                    style.append(f"border-{side}:{'2px' if edge.style in ('medium', 'thick') else '1px'} solid #333")
            value = cell.value
            if isinstance(value, str) and value.startswith("=") and results is not None:
                calculated = results.cell(row=r, column=c).value
                if calculated is None:  # never calculated (the file was not saved by Excel since)
                    style.append("color:#888;font-style:italic")
                else:
                    value = calculated
            if isinstance(value, (int, float)) and not isinstance(value, bool) and not (align and align.horizontal):
                style.append("text-align:right")
            rowspan, colspan = spans.get((r, c), (1, 1))
            attrs = (f' rowspan="{rowspan}"' if rowspan > 1 else "") + (f' colspan="{colspan}"' if colspan > 1 else "")
            text = html.escape(show_cell(value, 1000))
            cells.append(f'<td{attrs} style="{";".join(style)}">{text}</td>')
        rows.append(f'<tr style="height:{round(height * 1.33) if height else 20}px">{"".join(cells)}</tr>')
    header = "".join(f"<th>{get_column_letter(c)}</th>" for c in range(min_col, max_col + 1))
    return ("<!doctype html><html><head><meta charset='utf-8'><style>"
            "body{margin:0;background:#fff;font-family:Calibri,Carlito,Arial,sans-serif;font-size:15px}"
            "table{border-collapse:collapse;table-layout:fixed}"
            "td{border:1px solid #e1e1e1;padding:1px 4px;overflow:hidden;white-space:nowrap;vertical-align:bottom}"
            "th{background:#f3f3f3;border:1px solid #c8c8c8;font-weight:normal;color:#555;font-size:12px}"
            ".rh{width:34px}</style></head><body><table>"
            f'<colgroup><col style="width:34px">{"".join(cols)}</colgroup>'
            f"<tr><th></th>{header}</tr>{''.join(rows)}</table></body></html>")


def screenshot(page: str) -> bytes:
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        raise ToolError("Viewing a workbook without Excel needs the browser of screenshot_page: pip install "
                        "\"codeagent[browser]\", then playwright install chromium (or install Excel and xlwings).")
    from .browser import launch_browser

    with sync_playwright() as pw:
        browser, _ = launch_browser(pw)
        try:
            tab = browser.new_page(viewport={"width": 1400, "height": 900})
            tab.set_content(page)
            box = tab.eval_on_selector("table", "t => [t.scrollWidth, t.scrollHeight]")
            tab.set_viewport_size({"width": min(max(box[0] + 2, 200), 4000), "height": min(max(box[1] + 2, 100), 6000)})
            return tab.screenshot(full_page=False, type="png")
        finally:
            browser.close()


# --- format_excel --------------------------------------------------------------------------------------

STYLE_KEYS = ("bold", "italic", "font_color", "fill", "wrap", "align", "border", "number_format", "column_width")


def tool_format_excel(path: str, range: str, sheet: str | None = None, bold: bool | None = None,
                      italic: bool | None = None, font_color: str | None = None, fill: str | None = None,
                      wrap: bool | None = None, align: str | None = None, border: bool | None = None,
                      number_format: str | None = None, column_width=None, freeze: str | None = None,
                      autofilter: bool | None = None) -> str:
    if not state.excel_allow_format:
        raise ToolError("Formatting is not allowed in this application: write values only, the workbook's "
                        "layout stays as it is.")
    p = excel_path(path)
    if is_protected(p):
        raise ToolError(f"{path} is part of the coding agent's own files and cannot be modified.")
    name = rel_name(p)
    style = {"bold": bold, "italic": italic, "font_color": font_color, "fill": fill, "wrap": wrap, "align": align,
             "border": border, "number_format": number_format, "column_width": column_width,
             "freeze": freeze.upper().strip() if freeze else None, "autofilter": autofilter}
    if all(v is None for v in style.values()):
        raise ToolError("No formatting given (bold, fill, wrap, column_width, freeze, autofilter...).")
    for key in ("font_color", "fill"):
        if style[key] is not None and not COLOR.match(style[key]):
            raise ToolError(f"{key}: give a color as #RRGGBB, e.g. #DDEBF7.")
    if align is not None and align not in ("left", "center", "right"):
        raise ToolError("align: left, center or right.")
    if column_width is not None and column_width != "auto":
        try:
            if not 0 < float(column_width) <= 255:
                raise ValueError
        except (TypeError, ValueError):
            raise ToolError("column_width: 'auto' (fit the content) or a width from 1 to 255 (characters).")
    if style["freeze"] and not CELL.match(style["freeze"]):
        raise ToolError("freeze: the first cell NOT frozen, e.g. 'A2' to keep the first row in view.")
    ranges = [r.strip().upper().replace("$", "") for r in range.split(",") if r.strip()]
    if not ranges or not all(RANGE.match(r) for r in ranges):
        raise ToolError(f"Invalid range {range!r}: e.g. 'A1:G1', 'A:G' or 'A1:G1, A2:A40'.")

    backend = excel_backend()
    before = p.read_bytes()
    wb = load_workbook(p)
    if len(wb.sheetnames) > 1 and not sheet:
        raise ToolError(f"{name} has several sheets ({', '.join(wb.sheetnames)}): give 'sheet'. Nothing was changed.")
    sheet = sheet or wb.sheetnames[0]
    if sheet not in wb.sheetnames:
        raise ToolError(f"No sheet {sheet!r} in {name}. Sheets: {', '.join(wb.sheetnames)}")
    allowed = state.excel_edit_sheets.get(p.resolve())
    if allowed and sheet not in allowed:
        raise ToolError(f"Only these sheets of {name} may be changed: {', '.join(sorted(allowed))}.")

    described = [f"{k.replace('_', ' ')}: {v}" for k, v in style.items() if v is not None]
    lossy = lossy_features(p) if backend == "openpyxl" else []
    state.ui.panel(f"Format {name}, sheet {sheet}, {', '.join(ranges)}", described)
    if lossy:
        state.ui.warning(f"{name} contains {', '.join(lossy)}; saving with this tool removes or damages "
                         "them. A backup is kept, but check the file in Excel afterwards.")
    if state.auto_mode and not lossy:
        state.ui.status(f"(autonomous mode: formatting {name} without asking)")
    elif state.ui.confirm(f"Apply this formatting to {name}?") != "yes":
        feedback = state.ui.ask_text("Why not / what should change? (optional): ")
        raise ToolError(f"The user rejected the formatting; {name} was NOT changed."
                        + (f" User feedback: {feedback}" if feedback else ""))
    if p.read_bytes() != before:
        raise ToolError(f"{name} was changed on disk (probably saved in Excel) meanwhile; nothing was written. Try again.")
    backup = backups.save(p, before)
    if backend == "xlwings":
        from . import excel_xl
        skipped = excel_xl.apply_format(p, sheet, ranges, style)
    else:
        skipped = format_openpyxl(wb[sheet], ranges, style)
        save_openpyxl(wb, p, name)
    state.ui.success(f"Formatted {name} ({sheet} {', '.join(ranges)})")
    return (f"Formatted {display(p)}, sheet {sheet}, {', '.join(ranges)}: {', '.join(described)}."
            + (f" Not applied: {', '.join(skipped)}." if skipped else "")
            + f" Previous version saved to {backup}."
            + (f" Warning: {', '.join(lossy)} may have been lost." if lossy else "")
            + " Check the result with view_excel.")


def cells_of(ws, address: str):
    """Every cell of a range; a whole column or row stops at the sheet's used area."""
    from openpyxl.utils import range_boundaries

    min_col, min_row, max_col, max_row = range_boundaries(address)
    min_col, min_row = min_col or 1, min_row or 1
    max_col, max_row = max_col or ws.max_column, max_row or ws.max_row
    for row in ws.iter_rows(min_row=min_row, max_row=max(max_row, min_row), min_col=min_col, max_col=max(max_col, min_col)):
        yield from row


def format_openpyxl(ws, ranges: list[str], style: dict) -> list[str]:
    from copy import copy

    from openpyxl.styles import Border, PatternFill, Side
    from openpyxl.utils import get_column_letter, range_boundaries

    thin = Side(style="thin", color="000000")
    for address in ranges:
        for cell in cells_of(ws, address):
            if type(cell).__name__ == "MergedCell":
                continue
            if any(style[k] is not None for k in ("bold", "italic", "font_color")):
                font = copy(cell.font)
                if style["bold"] is not None:
                    font.b = style["bold"]
                if style["italic"] is not None:
                    font.i = style["italic"]
                if style["font_color"]:
                    font.color = "FF" + style["font_color"].lstrip("#").upper()
                cell.font = font
            if style["fill"]:
                color = "FF" + style["fill"].lstrip("#").upper()
                cell.fill = PatternFill(fill_type="solid", fgColor=color, bgColor=color)
            if style["wrap"] is not None or style["align"]:
                alignment = copy(cell.alignment)
                if style["wrap"] is not None:
                    alignment.wrap_text = style["wrap"]
                if style["align"]:
                    alignment.horizontal = style["align"]
                cell.alignment = alignment
            if style["border"] is not None:
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin) if style["border"] else Border()
            if style["number_format"]:
                cell.number_format = style["number_format"]
        if style["column_width"] is not None:
            min_col, _, max_col, _ = range_boundaries(address)
            for c in range(min_col or 1, (max_col or ws.max_column) + 1):
                letter = get_column_letter(c)
                if style["column_width"] == "auto":
                    longest = max((max(len(line) for line in str(v).splitlines() or [""])
                                   for (v,) in ws.iter_rows(min_col=c, max_col=c, values_only=True) if v is not None),
                                  default=8)
                    ws.column_dimensions[letter].width = min(max(longest * 1.1 + 2, 6), 80)
                else:
                    ws.column_dimensions[letter].width = float(style["column_width"])
    if style["autofilter"]:
        ws.auto_filter.ref = ranges[0] if ":" in ranges[0] else ws.dimensions
    if style["freeze"]:
        # The view is written whole, as Excel writes it: scrolled to the top, the frozen pane starting
        # at the freeze cell, the cursor there. openpyxl alone would keep the old scroll position (e.g.
        # topLeftCell A43 from a sheet left scrolled down) next to a pane starting at A2: Excel then
        # opens the sheet in a contradictory state.
        ws.freeze_panes = style["freeze"]
        ws.sheet_view.topLeftCell = "A1"
        for selection in ws.sheet_view.selection:
            selection.activeCell = selection.sqref = style["freeze"]
    return []
