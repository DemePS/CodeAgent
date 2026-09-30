"""Documents: read PDFs, read and edit Excel workbooks, view images."""

import base64
import io
import os
import re
from pathlib import Path

from openpyxl.utils import column_index_from_string, get_column_letter

from .. import backups, state
from ..common import (
    ToolError,
    display,
    is_protected,
    rel_name,
    resolve,
    resolve_readable,
    truncate,
)
from ..config import (
    EXCEL_BACKEND,
    EXCEL_CELL_CHARS,
    EXCEL_MAX_CELLS,
    EXCEL_MAX_CHANGES,
    IMAGE_TYPES,
    MAX_IMAGE_BYTES,
    MAX_TOOL_OUTPUT_CHARS,
    PDF_MAX_VISUAL_PAGES,
)

# --- Seeing pages and images --------------------------------------------------------------------------
# Claude reads images directly, so a screenshot gives it the page's text *and* its layout, spacing
# and colors -- better than OCR. Images go back as tool_result image blocks.

def image_block(data: bytes, media_type: str) -> dict:
    return {"type": "image", "source": {"type": "base64", "media_type": media_type,
                                        "data": base64.b64encode(data).decode("ascii")}}


# --- PDFs and Excel workbooks ----------------------------------------------------------------------

def parse_pages(spec: str | None, count: int) -> list[int]:
    """'2,4,10-12' -> [2, 4, 10, 11, 12] (1-based), validated against the page count."""
    if not spec:
        return list(range(1, count + 1))
    pages: list[int] = []
    for part in spec.replace(" ", "").split(","):
        match = re.fullmatch(r"(\d+)(?:-(\d+))?", part)
        if not match:
            raise ToolError(f"Invalid pages {spec!r}: use e.g. '3', '1-5' or '2,4,10-12'.")
        first, last = int(match.group(1)), int(match.group(2) or match.group(1))
        if not 1 <= first <= last <= count:
            raise ToolError(f"Pages {part} are outside the document (it has {count} page(s)).")
        pages.extend(p for p in range(first, last + 1) if p not in pages)
    return pages


SPREADSHEET_WORDS = re.compile(r"\.xls[xm]?\b|excel|spreadsheet|workbook|tableur|classeur", re.I)


def tool_read_pdf(path: str, pages: str | None = None, mode: str = "visual") -> list | str:
    # Spreadsheet first: know which fields are needed before reading documents.
    if SPREADSHEET_WORDS.search(state.turn["instruction"]) and not state.turn["excel_read"]:
        raise ToolError("This task involves a spreadsheet: open it with read_excel first, work out which "
                        "cells/fields must be filled (headers, units, formats, formula cells), list them, and "
                        "only then read the PDF pages that contain those fields.")
    try:
        from pypdf import PdfReader, PdfWriter
        from pypdf.errors import PdfReadError
    except ImportError:
        raise ToolError("pypdf is not installed in the agent's environment (uv sync).")
    p = resolve_readable(path)
    if not p.is_file():
        raise ToolError(f"File not found: {path}")
    try:
        reader = PdfReader(str(p))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ToolError(f"{path} is password-protected.")
        count = len(reader.pages)
    except PdfReadError as e:
        raise ToolError(f"{path} is not a readable PDF: {e}")
    selected = parse_pages(pages, count)
    label = f"{display(p)}: {count} page(s); showing page(s) {pages or ('1-' + str(count))}"

    if mode == "text":
        state.ui.status(f"[pdf] {rel_name(p)} text, {len(selected)} page(s)")
        parts = []
        for n in selected:
            text = (reader.pages[n - 1].extract_text() or "").strip()
            parts.append(f"--- page {n} ---\n{text or '(no text layer: a scan or an image -- use mode visual)'}")
        return truncate(label + "\n" + "\n".join(parts))

    if len(selected) > PDF_MAX_VISUAL_PAGES:
        raise ToolError(f"{path} has {count} pages; read at most {PDF_MAX_VISUAL_PAGES} at a time in visual mode "
                        f"(e.g. pages='1-{PDF_MAX_VISUAL_PAGES}'), or use mode='text' to skim all of it first.")
    if len(selected) == count:
        data = p.read_bytes()
    else:  # only the requested pages, as a smaller PDF
        writer = PdfWriter()
        for n in selected:
            writer.add_page(reader.pages[n - 1])
        buffer = io.BytesIO()
        writer.write(buffer)
        data = buffer.getvalue()
    if len(data) > 20 * 1024 * 1024:
        raise ToolError(f"The selected pages are {len(data):,} bytes (limit 20 MB); read fewer pages at a time.")
    state.ui.status(f"[pdf] {rel_name(p)} ({len(selected)} of {count} page(s), {len(data):,} bytes)")
    return [
        {"type": "text", "text": label + (" (page numbers inside the document below restart at 1)"
                                          if len(selected) != count else "")},
        {"type": "document", "title": display(p), "context": f"pages: {len(selected)} ({pages or 'all'} of {count})",
         "source": {"type": "base64", "media_type": "application/pdf",
                    "data": base64.b64encode(data).decode("ascii")}},
    ]


def excel_path(path: str, must_exist: bool = True, readable: bool = False) -> Path:
    """A workbook path: readable anywhere the agent may read, writable only in the workspace."""
    p = resolve_readable(path) if readable else resolve(path)
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        hint = " Save it as .xlsx in Excel first." if p.suffix.lower() == ".xls" else ""
        raise ToolError(f"{path}: only .xlsx and .xlsm workbooks are supported.{hint} For .csv use read_file/edit_file.")
    if must_exist and not p.is_file():
        raise ToolError(f"File not found: {path}")
    return p


# Workbooks read recently, kept while their file does not change: a large workbook takes seconds to
# load, and Claude reads the same one many times (sheet after sheet, range after range).
_READ_CACHE: dict[tuple[Path, bool], tuple[tuple[int, int], object]] = {}
_READ_CACHE_SIZE = 4


def read_workbook(p: Path, data_only: bool = False):
    """A workbook for reading only (never modified or saved): loaded once while the file is unchanged."""
    try:
        st = p.stat()
    except OSError as e:
        raise ToolError(f"{rel_name(p)} cannot be read: {e}")
    key, stamp = (p.resolve(), data_only), (st.st_mtime_ns, st.st_size)
    cached = _READ_CACHE.get(key)
    if cached and cached[0] == stamp:
        return cached[1]
    wb = load_workbook(p, data_only=data_only)
    # Each sheet's size as saved: reading cells (e.g. a range past the data) can add empty cells to a
    # loaded workbook and grow max_row, which must not change the sizes shown by later reads.
    wb.agent_sizes = {ws.title: (ws.max_row, ws.max_column) for ws in wb.worksheets}
    _READ_CACHE.pop(key, None)
    _READ_CACHE[key] = (stamp, wb)
    while len(_READ_CACHE) > _READ_CACHE_SIZE:
        _READ_CACHE.pop(next(iter(_READ_CACHE)))
    return wb


def sheet_size(wb, name: str) -> str:
    rows, cols = getattr(wb, "agent_sizes", {}).get(name) or (wb[name].max_row, wb[name].max_column)
    return f"{rows} rows x {cols} cols"


def load_workbook(p: Path, **options):
    try:
        import openpyxl
    except ImportError:
        raise ToolError("openpyxl is not installed in the agent's environment (uv sync).")
    try:
        return openpyxl.load_workbook(str(p), keep_vba=p.suffix.lower() == ".xlsm", **options)
    except PermissionError:
        raise ToolError(f"{rel_name(p)} is locked by another program (is it open in Excel?). Ask the user to close it.")
    except Exception as e:
        raise ToolError(f"{rel_name(p)} could not be opened as a workbook: {type(e).__name__}: {e}")


def show_cell(value, limit: int = EXCEL_CELL_CHARS) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    text = str(value)
    return text if len(text) <= limit else text[:limit] + f"... [{len(text) - limit} more characters: read this cell alone to see it whole]"


def excel_backend() -> str:
    """"xlwings" or "openpyxl": who changes, formats and renders workbooks (see excel_xl)."""
    from . import excel_xl

    if EXCEL_BACKEND == "openpyxl":
        return "openpyxl"
    if EXCEL_BACKEND == "xlwings":
        if not excel_xl.installed():
            raise ToolError("AGENT_EXCEL_BACKEND is xlwings but xlwings is not installed (pip install xlwings).")
        return "xlwings"
    return "xlwings" if excel_xl.available() else "openpyxl"


OVERVIEW_BUDGET = 12_000  # characters, whatever the number of sheets
OVERVIEW_VALUE_CHARS = 60  # a long cell value is cut in the overview


def workbook_overview(p: Path, wb, names: list[str]) -> str:
    """A workbook with several sheets, read without naming one: every sheet's name and size, then
    the first rows of each while the size budget allows (fewer rows when there are many sheets), so
    that Claude picks the sheet(s) to read in full instead of reading the first one blindly."""
    allowed = state.excel_edit_sheets.get(p.resolve())
    rows_each = 6 if len(names) <= 10 else 3 if len(names) <= 30 else 1
    lines = [f"{display(p)} -- {len(names)} sheets (overview; read one in full with sheet=..., or part of it with range=...)"]
    if allowed:
        lines.append("sheets to fill (the person chose them; others cannot be changed): " + ", ".join(sorted(allowed)))
    sizes = getattr(wb, "agent_sizes", {})
    lines.append("sheets: " + "; ".join(f"{n} ({'x'.join(map(str, sizes.get(n) or (wb[n].max_row, wb[n].max_column)))})" for n in names))
    lines.append(f"--- first {'row' if rows_each == 1 else f'{rows_each} rows'} of each sheet ---")
    size = sum(len(line) + 1 for line in lines)
    for number, n in enumerate(names):
        preview = []
        saved_rows = (sizes.get(n) or (wb[n].max_row, 0))[0]
        for row in wb[n].iter_rows(max_row=min(rows_each, saved_rows)):
            parts = []
            for c in row:
                if c.value is None or not hasattr(c, "coordinate"):
                    continue
                value = show_cell(c.value)
                parts.append(f"{c.coordinate}={value[:OVERVIEW_VALUE_CHARS]}{'…' if len(value) > OVERVIEW_VALUE_CHARS else ''}")
            if parts:
                preview.append(" | ".join(parts[:20]))
        block = [f"[{n}]"] + (preview or ["(empty in its first rows)"])
        block_size = sum(len(line) + 1 for line in block)
        if size + block_size > OVERVIEW_BUDGET:
            lines.append(f"... no preview for the {len(names) - number} remaining sheet(s) (size limit): read them "
                         "with sheet=... if their name suggests they matter")
            break
        lines += block
        size += block_size
    state.ui.status(f"[excel] {rel_name(p)}: overview of {len(names)} sheets")
    return truncate("\n".join(lines))


def tool_read_excel(path: str, sheet: str | None = None, range: str | None = None) -> str:
    state.turn["excel_read"] = True  # any attempt counts: the workbook may not exist yet (to be created)
    p = excel_path(path, readable=True)
    from . import excel_lock
    lock_note = excel_lock.try_hold(p)  # nobody else changes it until the instruction ends
    formulas = read_workbook(p)  # formulas as written
    names = formulas.sheetnames
    if sheet is None and range is None and len(names) > 1:
        return workbook_overview(p, formulas, names) + (f"\n{lock_note}" if lock_note else "")
    ws_name = sheet or names[0]
    if ws_name not in names:
        raise ToolError(f"No sheet {ws_name!r}. Sheets: {', '.join(names)}")
    ws = formulas[ws_name]
    values_sheet = []  # the values Excel calculated last time it saved: loaded only if a formula is shown

    def calculated(coordinate: str):
        if not values_sheet:
            values_sheet.append(read_workbook(p, data_only=True)[ws_name])
        return values_sheet[0][coordinate].value

    # The other sheets are listed only for a small workbook: with many sheets, the overview lists them
    # once, instead of every read repeating the list.
    others = ("sheets: " + "; ".join(f"{n} ({sheet_size(formulas, n)})" for n in names)
              if len(names) <= 10 else f"one of {len(names)} sheets (read_excel without sheet lists them)")
    lines = [f"{display(p)} -- sheet {ws_name} ({sheet_size(formulas, ws_name)}); {others}"]
    if lock_note:
        lines.append(lock_note)
    try:
        cells = ws[range] if range else ws.iter_rows()
    except ValueError:
        raise ToolError(f"Invalid range {range!r}; use e.g. 'A1:F40'.")
    single = bool(range) and not isinstance(cells, tuple)
    if single:  # a single cell: shown whole, however long
        rows = ((cells,),)
    elif range and cells and not isinstance(cells[0], tuple):  # a single column such as "A:A"
        rows = tuple((c,) for c in cells)
    else:
        rows = cells
    if ws.merged_cells.ranges:
        lines.append("merged: " + ", ".join(str(r) for r in list(ws.merged_cells.ranges)[:50]))
    lines += sheet_objects(ws)
    lines.append(f"--- sheet {ws_name}" + (f" range {range}" if range else "") + " ---")
    shown = size = 0

    def rows_seen_label(row) -> str:
        return str(next((c.row for c in row if hasattr(c, "row")), "?"))

    for row in rows:
        parts = []
        for c in row:
            if c.value is None or not hasattr(c, "coordinate"):
                continue
            limit = MAX_TOOL_OUTPUT_CHARS if single else EXCEL_CELL_CHARS
            if isinstance(c.value, str) and c.value.startswith("="):
                cached = calculated(c.coordinate)
                parts.append(f"{c.coordinate}={c.value} -> {show_cell(cached, limit) if cached is not None else '(not calculated)'}")
            else:
                fmt = f" [{c.number_format}]" if c.number_format not in ("General", None) else ""
                parts.append(f"{c.coordinate}={show_cell(c.value, limit)}{fmt}")
        if parts:
            lines.append(" | ".join(parts))
            shown += len(parts)
            size += len(lines[-1]) + 1
            if shown >= EXCEL_MAX_CELLS or size > MAX_TOOL_OUTPUT_CHARS - 1000:
                last = rows_seen_label(row)
                lines.append(f"... stopped at row {last} after {shown} cells; read the rest with a range "
                             f"starting below row {last}")
                break
    if shown == 0:
        lines.append("(no values)")
    state.ui.status(f"[excel] {rel_name(p)} sheet {ws_name}{' ' + range if range else ''}: {shown} cell(s)")
    return truncate("\n".join(lines))


def anchor_cell(obj) -> str:
    """The top-left cell a chart or picture is placed on, e.g. 'D2' ('?' if unknown)."""
    try:
        marker = obj.anchor._from
        return f"{get_column_letter(marker.col + 1)}{marker.row + 1}"
    except Exception:
        return "?"


def chart_title(chart) -> str:
    try:
        return "".join(r.t for p in chart.title.tx.rich.p for r in (p.r or []))
    except Exception:
        return ""


def sheet_objects(ws) -> list[str]:
    """What the sheet holds besides cells: charts (type, title, place, the cells they plot), pictures,
    Excel tables and pivot tables -- so they are known before anything is changed."""
    lines = []
    for chart in getattr(ws, "_charts", []):
        refs = []
        for series in getattr(chart, "series", [])[:6]:
            for part in (getattr(series, "cat", None), getattr(series, "val", None)):
                ref = getattr(getattr(part, "numRef", None), "f", None) or getattr(getattr(part, "strRef", None), "f", None)
                if ref and ref not in refs:
                    refs.append(ref)
        title = chart_title(chart)
        lines.append(f"chart: {type(chart).__name__.replace('Chart', '').lower() or 'chart'} chart"
                     + (f" '{title}'" if title else "") + f" at {anchor_cell(chart)}"
                     + (f", plotting {', '.join(refs)}" if refs else ""))
    for image in getattr(ws, "_images", []):
        lines.append(f"picture at {anchor_cell(image)}")
    for name, table in getattr(ws, "tables", {}).items():
        lines.append(f"Excel table '{name}': {getattr(table, 'ref', table)}")
    for pivot in getattr(ws, "_pivots", []):
        lines.append(f"pivot table '{getattr(pivot, 'name', '')}' at {getattr(getattr(pivot, 'location', None), 'ref', '?')}")
    return lines


def lossy_features(p: Path) -> list[str]:
    """Parts of an .xlsx that openpyxl drops or damages when it saves the file."""
    import zipfile
    try:
        names = zipfile.ZipFile(p).namelist()
    except (zipfile.BadZipFile, OSError):
        return []
    found = []
    for prefix, label in (("xl/charts/", "charts"), ("xl/media/", "images"), ("xl/pivotTables/", "pivot tables"),
                          ("xl/slicers/", "slicers"), ("xl/externalLinks/", "links to other workbooks"),
                          ("xl/threadedComments/", "threaded comments"), ("xl/ctrlProps/", "form controls")):
        if any(n.startswith(prefix) for n in names):
            found.append(label)
    return found


def tool_edit_excel(path: str, changes: list, create_sheets: list | None = None) -> str:
    import datetime
    p = excel_path(path, must_exist=False)
    if is_protected(p):
        raise ToolError(f"{path} is part of the coding agent's own files and cannot be modified.")
    if not changes and not create_sheets:
        raise ToolError("No changes given.")
    if len(changes) > EXCEL_MAX_CHANGES:
        raise ToolError(f"At most {EXCEL_MAX_CHANGES} cells per call; split the changes.")
    existed = p.exists()
    backend = excel_backend()
    from . import excel_lock
    if existed:
        excel_lock.hold(p)  # refused if another program has it open
    before = p.read_bytes() if existed else None
    if existed:  # checked (and, with openpyxl, changed) on openpyxl's copy; xlwings has Excel apply it
        wb = load_workbook(p)
    else:
        import openpyxl
        wb = openpyxl.Workbook()
    name = rel_name(p)
    allowed = state.excel_edit_sheets.get(p.resolve())
    if allowed:
        others = sorted({s for s in create_sheets or []} | {c.get("sheet") or wb.sheetnames[0] for c in changes
                                                              if isinstance(c, dict)} - allowed)
        if others:
            raise ToolError(f"Only these sheets of {name} may be filled: {', '.join(sorted(allowed))} (the person chose "
                            f"them). Not: {', '.join(others)}. If a value really belongs there, ask the person.")

    for sheet_name in create_sheets or []:
        if sheet_name in wb.sheetnames:
            raise ToolError(f"Sheet {sheet_name!r} already exists.")
        wb.create_sheet(sheet_name)
    if len(wb.sheetnames) > 1 and any(isinstance(c, dict) and not c.get("sheet") for c in changes):
        raise ToolError(f"{name} has several sheets ({', '.join(wb.sheetnames)}): give 'sheet' for every change. "
                        "Nothing was changed.")
    widths = state.excel_max_columns.get(p.resolve(), {})
    rows, writes = [], []
    for change in changes:
        sheet_name = change.get("sheet") or wb.sheetnames[0]
        if sheet_name not in wb.sheetnames:
            raise ToolError(f"No sheet {sheet_name!r} in {name}. Sheets: {', '.join(wb.sheetnames)} "
                            "(add it with create_sheets).")
        ws = wb[sheet_name]
        coord = str(change.get("cell", "")).upper().strip()
        if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", coord):
            raise ToolError(f"Invalid cell {change.get('cell')!r}; use an address such as 'B7'.")
        if sheet_name in widths and column_index_from_string(re.match(r"[A-Z]+", coord).group()) > widths[sheet_name]:
            raise ToolError(f"{sheet_name}!{coord} is outside the sheet's columns (A to {get_column_letter(widths[sheet_name])}): "
                            "do not add columns or helper cells. Nothing was changed.")
        cell = ws[coord]
        if type(cell).__name__ == "MergedCell":
            raise ToolError(f"{sheet_name}!{coord} is inside a merged range; write to its top-left cell instead.")
        value = change.get("value")
        if change.get("as_date") and isinstance(value, str):
            try:
                value = datetime.date.fromisoformat(value[:10]) if len(value) <= 10 else datetime.datetime.fromisoformat(value)
            except ValueError:
                raise ToolError(f"{sheet_name}!{coord}: {value!r} is not an ISO date (YYYY-MM-DD).")
        if isinstance(value, (dict, list)):
            raise ToolError(f"{sheet_name}!{coord}: a cell value must be a number, text, boolean or null.")
        old = cell.value
        if state.excel_protect_formulas:
            if isinstance(value, str) and value.startswith("="):
                raise ToolError(f"{sheet_name}!{coord}: formulas may not be written here; write values only. Nothing was changed.")
            if isinstance(old, str) and old.startswith("=") and value != old:
                raise ToolError(f"{sheet_name}!{coord} holds a formula ({old[:60]}), which may not be changed or cleared. "
                                "Nothing was changed.")
        cell.value = value
        if change.get("number_format"):
            cell.number_format = change["number_format"]
        writes.append((sheet_name, coord, value, change.get("number_format")))
        if show_cell(old) != show_cell(value) or change.get("number_format"):
            rows.append((f"{sheet_name}!{coord}", show_cell(old), show_cell(value), change.get("number_format")))

    lossy = lossy_features(p) if existed and backend == "openpyxl" else []  # Excel keeps everything
    state.ui.cell_changes(f"{'Modify' if existed else 'Create'} workbook {name} ({len(rows)} cell(s)"
                          f"{', new sheets: ' + ', '.join(create_sheets) if create_sheets else ''})",
                          rows[:200], max(0, len(rows) - 200))
    if lossy:
        state.ui.warning(f"{name} contains {', '.join(lossy)}; saving with this tool removes or damages "
                         "them. A backup is kept, but check the file in Excel afterwards.")
    question = f"Apply these cell changes to {name}?" if existed else f"Create {name}?"
    if state.auto_mode and not lossy:
        state.ui.status(f"(autonomous mode: {'modifying' if existed else 'creating'} {name} without asking)")
    else:  # lossy saves always ask, even in autonomous mode
        if state.auto_mode:
            state.ui.status("(autonomous mode: this save can lose content, so it needs your approval)")
        if state.ui.confirm(question) != "yes":
            feedback = state.ui.ask_text(f"Why not / what should change in {name}? (optional): ")
            state.ui.failure(f"{name} was not {'modified' if existed else 'created'}")
            raise ToolError(f"The user rejected this change; {name} was NOT {'modified' if existed else 'created'}."
                            + (f" User feedback: {feedback}" if feedback else ""))

    if (p.read_bytes() if p.exists() else None) != before:
        state.ui.failure(f"{name} changed on disk while waiting for approval -- not written")
        raise ToolError(f"{name} was changed on disk (probably saved in Excel) after the diff was shown; nothing "
                        "was written. Read it again and redo the changes.")
    backup = None
    if existed:  # a copy of the previous version, outside the project (see backups.py)
        backup = backups.save(p, before)
    p.parent.mkdir(parents=True, exist_ok=True)
    if backend == "xlwings":
        from . import excel_xl
        excel_xl.apply_changes(p, existed, list(create_sheets or []), writes)
    else:
        save_openpyxl(wb, p, name)
    if not existed:
        excel_lock.try_hold(p)  # a new workbook is the agent's too until the instruction ends
    state.ui.success(f"{'Modified' if existed else 'Created'} {name} ({len(rows)} cell(s))")
    return (f"{'Modified' if existed else 'Created'} {display(p)}: {len(rows)} cell(s) changed."
            + (f" Previous version saved to {backup}." if backup else "")
            + (f" Warning: {', '.join(lossy)} may have been lost." if lossy else "")
            + (" Excel recalculated the formulas." if backend == "xlwings" else
               " Formulas are recalculated when the file is opened in Excel." if any(
                isinstance(r[2], str) and r[2].startswith("=") for r in rows) else ""))


def save_openpyxl(wb, p: Path, name: str) -> None:
    from . import excel_lock

    tmp = p.with_name(f".{p.name}.agent-tmp")
    try:
        wb.save(str(tmp))
        if not excel_lock.write(p, tmp.read_bytes()):  # held: written through the lock
            os.replace(tmp, p)
    except PermissionError:
        raise ToolError(f"{name} could not be written: it is locked (open in Excel?). Ask the user to close it.")
    finally:
        if tmp.exists():
            tmp.unlink()


def tool_view_image(path: str) -> list:
    p = resolve_readable(path)
    if not p.is_file():
        raise ToolError(f"File not found: {path}")
    media_type = IMAGE_TYPES.get(p.suffix.lower())
    if media_type is None:
        raise ToolError(f"Not a supported image ({', '.join(IMAGE_TYPES)}). For SVG, read it with read_file.")
    data = p.read_bytes()
    if len(data) > MAX_IMAGE_BYTES:
        raise ToolError(f"{path} is {len(data):,} bytes; images must be under {MAX_IMAGE_BYTES:,}. "
                        "Ask the user for a smaller export.")
    state.ui.status(f"[image] {rel_name(p)} ({len(data):,} bytes)")
    return [{"type": "text", "text": f"{display(p)} ({len(data):,} bytes):"}, image_block(data, media_type)]




def tool_restore_backup(path: str, version: str | None = None) -> str:
    """List a workbook's previous versions, or put one back (the person approves, like any change)."""
    from ..cleanup import BACKUP_DAYS, days

    p = excel_path(path, must_exist=False)
    if is_protected(p):
        raise ToolError(f"{path} is part of the coding agent's own files and cannot be modified.")
    name = rel_name(p)
    found = backups.versions(p.name)
    if not found:
        raise ToolError(f"No previous version of {name}: the agent keeps a copy before each change it saves, "
                        f"for {days('AGENT_BACKUP_DAYS', BACKUP_DAYS)} days, and it has none for this workbook.")
    if version is None:
        lines = [f"- version={v['id']!r}: as it was before the change of {backups.moment(v['id']):%a %d %b %Y %H:%M:%S}"
                 for v in found]
        return (f"Previous versions of {name}, newest first (each is the workbook as it was just before one of "
                "your saved changes):\n" + "\n".join(lines)
                + "\nTo put one back, call restore_backup again with its version. To undo a whole job, choose "
                  "the oldest version from that job.")
    if version not in {v["id"] for v in found}:
        raise ToolError(f"{version!r} is not a previous version of {name}; call restore_backup without a version to list them.")
    when = f"{backups.moment(version):%a %d %b %Y at %H:%M}"
    state.ui.panel(f"Restore {name}", [f"Put back {name} as it was before the change of {when}.",
                                       "The current version is kept among the previous versions."])
    if state.auto_mode:
        state.ui.status(f"(autonomous mode: restoring {name} without asking)")
    elif state.ui.confirm(f"Restore {name} to its version from before the change of {when}?") != "yes":
        feedback = state.ui.ask_text("Why not? (optional): ")
        raise ToolError(f"The user rejected the restore; {name} was NOT changed." + (f" User feedback: {feedback}" if feedback else ""))
    from . import excel_lock
    entry = excel_lock.hold(p)
    try:
        if entry is not None and entry.file is not None:  # openpyxl: through the lock
            message = backups.restore(p, version, write=lambda data: excel_lock.write(p, data))
        else:
            if entry is not None:  # xlwings: the agent's Excel lets it go for the swap, then takes it again
                excel_lock.release(p)
            message = backups.restore(p, version)
            if entry is not None:
                excel_lock.try_hold(p)
    except (ValueError, PermissionError) as e:
        raise ToolError(str(e))
    state.ui.success(f"Restored {name} (version from before {when})")
    return message
