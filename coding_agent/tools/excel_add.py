"""add_chart, add_table, add_pivot_table: add a chart, an Excel table or a pivot table to a workbook.

Each tool takes a fixed set of fields, checked here before anything is shown; the person approves
what will be added (like edit_excel and format_excel), a backup is kept, then the workbook is saved:
- with Excel (xlwings backend), Excel adds the object and saves -- nothing else in the file is touched;
- with openpyxl, charts and tables are written by openpyxl (the save rewrites the file: the person is
  warned first when that can damage what is already in it). Pivot tables need Excel: openpyxl cannot
  build one, and Excel alone computes it.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl.utils import get_column_letter, range_boundaries

from .. import backups, state
from ..common import ToolError, display, is_protected, rel_name
from .documents import excel_backend, excel_path, load_workbook, lossy_features, save_openpyxl

CHART_TYPES = ("column", "bar", "line", "pie", "area", "scatter")
SUMMARIES = ("sum", "count", "average", "min", "max")
TABLE_STYLE = "TableStyleMedium2"
CELL = re.compile(r"^[A-Z]{1,3}[1-9][0-9]{0,6}$")
BLOCK = re.compile(r"^[A-Z]{1,3}[1-9][0-9]{0,6}:[A-Z]{1,3}[1-9][0-9]{0,6}$")
# An Excel table or pivot table name: a letter or underscore first, then letters, digits, _ or .;
# never something Excel reads as a cell (A1, R1C1).
NAME = re.compile(r"^[A-Za-z_][A-Za-z0-9_.]{0,254}$")
CELL_LIKE = re.compile(r"^([A-Za-z]{1,3}[0-9]+|[Rr][0-9]*[Cc][0-9]*|[RrCc])$")


def tool_add_chart(path: str, source: str, chart_type: str, sheet: str | None = None, title: str | None = None,
                   anchor: str | None = None) -> str:
    return add(path, "chart", source, sheet, chart_type=chart_type, title=title, anchor=anchor)


def tool_add_table(path: str, source: str, sheet: str | None = None, name: str | None = None) -> str:
    return add(path, "table", source, sheet, name=name)


def tool_add_pivot_table(path: str, source: str, rows: list, values: list, sheet: str | None = None,
                         columns: list | None = None, target_sheet: str | None = None, anchor: str | None = None,
                         name: str | None = None) -> str:
    return add(path, "pivot_table", source, sheet, rows=rows, values=values, columns=columns,
               target_sheet=target_sheet, anchor=anchor, name=name)


def add(path: str, kind: str, source: str, sheet: str | None = None, chart_type: str | None = None,
        title: str | None = None, anchor: str | None = None, name: str | None = None, rows: list | None = None,
        columns: list | None = None, values: list | None = None, target_sheet: str | None = None) -> str:
    """What the three tools share: checks, the person's approval, a backup, then the save."""
    if not state.excel_allow_format:
        raise ToolError("Adding charts or tables is not allowed in this application: write values only, the "
                        "workbook's layout stays as it is.")
    p = excel_path(path)
    if is_protected(p):
        raise ToolError(f"{path} is part of the coding agent's own files and cannot be modified.")
    book_name = rel_name(p)
    source = (source or "").strip().upper().replace("$", "")
    if not BLOCK.match(source):
        raise ToolError(f"source {source!r}: the cells holding the data with their header row, e.g. 'A1:D20'.")
    min_col, min_row, max_col, max_row = range_boundaries(source)
    if max_row - min_row < 1:
        raise ToolError(f"source {source}: give the header row and at least one row of data.")
    anchor = anchor.strip().upper().replace("$", "") if anchor else None
    if anchor and not CELL.match(anchor):
        raise ToolError(f"anchor {anchor!r}: one cell, e.g. 'F2'.")

    backend = excel_backend()
    before = p.read_bytes()
    wb = load_workbook(p)
    if len(wb.sheetnames) > 1 and not sheet:
        raise ToolError(f"{book_name} has several sheets ({', '.join(wb.sheetnames)}): give 'sheet' (the one holding "
                        "the data). Nothing was changed.")
    sheet = sheet or wb.sheetnames[0]
    if sheet not in wb.sheetnames:
        raise ToolError(f"No sheet {sheet!r} in {book_name}. Sheets: {', '.join(wb.sheetnames)}")
    ws = wb[sheet]
    headers = [ws.cell(row=min_row, column=c).value for c in range(min_col, max_col + 1)]

    if kind == "chart":
        spec = chart_spec(chart_type, title, anchor or default_anchor(max_col, min_row), headers, source)
        written_on = spec["sheet"] = sheet
    elif kind == "table":
        spec = table_spec(wb, ws, name, source, headers)
        written_on = sheet
    else:
        if backend != "xlwings":
            raise ToolError("Pivot tables need Excel (the xlwings backend): openpyxl cannot build them. Install "
                            "Excel and the excel extra (uv sync --extra excel), or make a summary table with "
                            "formulas (SUMIFS, COUNTIFS) instead. Nothing was changed.")
        spec = pivot_spec(wb, name, rows, columns, values, headers, target_sheet, anchor)
        written_on = spec["target_sheet"]
    allowed = state.excel_edit_sheets.get(p.resolve())
    if allowed and written_on not in allowed:
        raise ToolError(f"Only these sheets of {book_name} may be changed: {', '.join(sorted(allowed))}.")

    description = describe(kind, spec, sheet, source)
    lossy = lossy_features(p) if backend == "openpyxl" else []
    state.ui.panel(f"Add a {kind.replace('_', ' ')} to {book_name}", description)
    if lossy:
        state.ui.warning(f"{book_name} contains {', '.join(lossy)}; saving with this tool removes or damages "
                         "them. A backup is kept, but check the file in Excel afterwards.")
    if state.auto_mode and not lossy:
        state.ui.status(f"(autonomous mode: adding a {kind.replace('_', ' ')} to {book_name} without asking)")
    else:  # lossy saves always ask, even in autonomous mode
        if state.auto_mode:
            state.ui.status("(autonomous mode: this save can lose content, so it needs your approval)")
        if state.ui.confirm(f"Add this {kind.replace('_', ' ')} to {book_name}?") != "yes":
            feedback = state.ui.ask_text("Why not / what should change? (optional): ")
            raise ToolError(f"The user rejected the {kind.replace('_', ' ')}; {book_name} was NOT changed."
                            + (f" User feedback: {feedback}" if feedback else ""))
    if p.read_bytes() != before:
        raise ToolError(f"{book_name} was changed on disk (probably saved in Excel) meanwhile; nothing was written. "
                        "Try again.")
    backup = backups.save(p, before)
    skipped: list[str] = []
    if backend == "xlwings":
        from . import excel_xl
        if kind == "chart":
            skipped = excel_xl.add_chart(p, sheet, source, spec)
        elif kind == "table":
            excel_xl.add_table(p, sheet, source, spec)
        else:
            excel_xl.add_pivot_table(p, sheet, source, spec)
    else:
        if kind == "chart":
            chart_openpyxl(ws, source, spec)
        else:
            table_openpyxl(ws, source, spec)
        save_openpyxl(wb, p, book_name)
    state.ui.success(f"Added a {kind.replace('_', ' ')} to {book_name}")
    return (f"Added to {display(p)}: " + "; ".join(description) + "."
            + (f" Not applied: {', '.join(skipped)}." if skipped else "")
            + f" Previous version saved to {backup}."
            + (f" Warning: {', '.join(lossy)} may have been lost." if lossy else "")
            + " Check the result with view_excel" + (" (charts are drawn only when Excel is available)."
                                                     if kind == "chart" and backend != "xlwings" else "."))


def default_anchor(max_col: int, min_row: int) -> str:
    """Beside the data: one empty column to its right, level with its header."""
    return f"{get_column_letter(max_col + 2)}{min_row}"


def header_names(headers: list, source: str) -> list[str]:
    names = [str(h).strip() if h is not None else "" for h in headers]
    if not all(names):
        raise ToolError(f"Every column of {source} needs a header in its first row (empty: column "
                        f"{names.index('') + 1}). Nothing was changed.")
    lowered = [n.lower() for n in names]
    duplicates = sorted({n for n in names if lowered.count(n.lower()) > 1})
    if duplicates:
        raise ToolError(f"The headers of {source} must differ: {', '.join(duplicates)} appear(s) twice. Nothing was changed.")
    return names


def chart_spec(chart_type: str | None, title: str | None, anchor: str, headers: list, source: str) -> dict:
    if chart_type not in CHART_TYPES:
        raise ToolError(f"chart_type: one of {', '.join(CHART_TYPES)}.")
    if len(headers) < 2:
        raise ToolError(f"source {source}: the first column holds the labels (or the x values of a scatter "
                        "chart), the next ones the numbers -- give at least two columns.")
    if chart_type == "pie" and len(headers) != 2:
        raise ToolError(f"A pie chart shows one series: give two columns (labels, then numbers), not {len(headers)}.")
    series = [str(h) if h is not None else f"column {i + 2}" for i, h in enumerate(headers[1:])]
    return {"chart_type": chart_type, "title": (title or "").strip() or None, "anchor": anchor, "series": series}


def existing_names(wb) -> set[str]:
    """Table, pivot table and defined names already in the workbook (lowercase: Excel ignores case)."""
    names = {n.lower() for ws in wb.worksheets for n in getattr(ws, "tables", {})}
    names |= {getattr(pivot, "name", "").lower() for ws in wb.worksheets for pivot in getattr(ws, "_pivots", [])}
    try:
        names |= {n.lower() for n in wb.defined_names}
    except TypeError:  # older openpyxl: a list of DefinedName
        names |= {d.name.lower() for d in wb.defined_names.definedName}
    return names


def new_name(wb, wanted: str | None, prefix: str) -> str:
    taken = existing_names(wb)
    if wanted:
        wanted = wanted.strip()
        if not NAME.match(wanted) or CELL_LIKE.match(wanted):
            raise ToolError(f"name {wanted!r}: letters, digits and _ only, starting with a letter, no spaces, and not "
                            "a cell address (e.g. 'Expenses', 'Sales_2026').")
        if wanted.lower() in taken:
            raise ToolError(f"name {wanted!r} is already used in this workbook; choose another.")
        return wanted
    number = 1
    while f"{prefix}{number}".lower() in taken:
        number += 1
    return f"{prefix}{number}"


def overlaps(a: str, b: str) -> bool:
    a1, a2, a3, a4 = range_boundaries(a)
    b1, b2, b3, b4 = range_boundaries(b)
    return a1 <= b3 and b1 <= a3 and a2 <= b4 and b2 <= a4


def table_spec(wb, ws, name: str | None, source: str, headers: list) -> dict:
    header_names(headers, source)
    for other, table in getattr(ws, "tables", {}).items():
        if overlaps(source, getattr(table, "ref", table)):
            raise ToolError(f"{source} overlaps the Excel table '{other}' ({getattr(table, 'ref', table)}). Nothing was changed.")
    for merged in ws.merged_cells.ranges:
        if overlaps(source, str(merged)):
            raise ToolError(f"{source} contains merged cells ({merged}); an Excel table cannot. Nothing was changed.")
    return {"name": new_name(wb, name, "Table"), "style": TABLE_STYLE}


def pivot_spec(wb, name: str | None, rows: list | None, columns: list | None, values: list | None, headers: list,
               target_sheet: str | None, anchor: str | None) -> dict:
    fields = header_names(headers, "the source")
    by_lower = {f.lower(): f for f in fields}

    def field(value, role: str) -> str:
        if not isinstance(value, str) or value.strip().lower() not in by_lower:
            raise ToolError(f"{role}: {value!r} is not a header of the source. Headers: {', '.join(fields)}.")
        return by_lower[value.strip().lower()]

    rows = [field(f, "rows") for f in rows or []]
    columns = [field(f, "columns") for f in columns or []]
    if not rows:
        raise ToolError("rows: the header(s) to group by, e.g. ['Supplier'].")
    if set(rows) & set(columns):
        raise ToolError(f"{', '.join(sorted(set(rows) & set(columns)))}: a field is either in rows or in columns, not both.")
    data = []
    for item in values or []:
        if isinstance(item, str):
            item = {"field": item}
        if not isinstance(item, dict):
            raise ToolError("values: a list of {'field': header, 'summary': 'sum'}.")
        summary = (item.get("summary") or "sum").lower()
        if summary not in SUMMARIES:
            raise ToolError(f"values: summary is one of {', '.join(SUMMARIES)}, not {summary!r}.")
        data.append((field(item.get("field"), "values"), summary))
    if not data:
        raise ToolError("values: what to total, e.g. [{'field': 'Amount', 'summary': 'sum'}].")
    if target_sheet is None:
        number, target_sheet = 1, "Pivot"
        while target_sheet.lower() in {s.lower() for s in wb.sheetnames}:
            number += 1
            target_sheet = f"Pivot {number}"
        new_sheet = True
    else:
        target_sheet = target_sheet.strip()
        if not target_sheet or len(target_sheet) > 31 or any(ch in target_sheet for ch in "[]:*?/\\"):
            raise ToolError("target_sheet: a sheet name of 1 to 31 characters, without [ ] : * ? / \\.")
        new_sheet = target_sheet not in wb.sheetnames
        if not new_sheet and anchor is None:
            raise ToolError(f"Sheet {target_sheet!r} exists: give anchor, an empty cell where the pivot table starts.")
        if not new_sheet and wb[target_sheet][anchor].value is not None:
            raise ToolError(f"{target_sheet}!{anchor} is not empty; the pivot table needs empty cells. Nothing was changed.")
    return {"name": new_name(wb, name, "Pivot"), "rows": rows, "columns": columns, "values": data,
            "target_sheet": target_sheet, "new_sheet": new_sheet, "anchor": anchor or "A3"}


def describe(kind: str, spec: dict, sheet: str, source: str) -> list[str]:
    if kind == "chart":
        return [f"{spec['chart_type']} chart" + (f" '{spec['title']}'" if spec["title"] else ""),
                f"data: {sheet}!{source} (labels from its first column; series: {', '.join(spec['series'])})",
                f"placed at {sheet}!{spec['anchor']}"]
    if kind == "table":
        return [f"Excel table '{spec['name']}' on {sheet}!{source} (header row, filter buttons, banded rows)"]
    where = f"{'new sheet ' if spec['new_sheet'] else ''}{spec['target_sheet']}!{spec['anchor']}"
    return [f"pivot table '{spec['name']}' from {sheet}!{source}, at {where}",
            "rows: " + ", ".join(spec["rows"]),
            *(["columns: " + ", ".join(spec["columns"])] if spec["columns"] else []),
            "values: " + ", ".join(f"{summary} of {f}" for f, summary in spec["values"])]


# --- openpyxl -----------------------------------------------------------------------------------------

def chart_openpyxl(ws, source: str, spec: dict) -> None:
    from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference, ScatterChart, Series

    min_col, min_row, max_col, max_row = range_boundaries(source)
    labels = Reference(ws, min_col=min_col, min_row=min_row + 1, max_row=max_row)
    kind = spec["chart_type"]
    if kind == "scatter":
        chart = ScatterChart()
        chart.style = 13
        for col in range(min_col + 1, max_col + 1):
            series = Series(Reference(ws, min_col=col, min_row=min_row, max_row=max_row), labels, title_from_data=True)
            series.marker.symbol = "circle"
            series.graphicalProperties.line.noFill = True
            chart.series.append(series)
    else:
        chart = {"column": BarChart, "bar": BarChart, "line": LineChart, "pie": PieChart, "area": AreaChart}[kind]()
        if kind in ("column", "bar"):
            chart.type = "col" if kind == "column" else "bar"
        chart.add_data(Reference(ws, min_col=min_col + 1, max_col=max_col, min_row=min_row, max_row=max_row),
                       titles_from_data=True)
        chart.set_categories(labels)
    if spec["title"]:
        chart.title = spec["title"]
    ws.add_chart(chart, spec["anchor"])


def table_openpyxl(ws, source: str, spec: dict) -> None:
    from openpyxl.worksheet.table import Table, TableStyleInfo

    table = Table(displayName=spec["name"], ref=source)
    table.tableStyleInfo = TableStyleInfo(name=spec["style"], showRowStripes=True)
    ws.add_table(table)
