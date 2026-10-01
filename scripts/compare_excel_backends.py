"""Compare the two Excel backends of the agent -- openpyxl (rewrites the file) and xlwings (Excel
itself does the work) -- by running the same tools on the same workbooks, then checking the files.

    uv run --extra excel --extra browser python scripts/compare_excel_backends.py            # sample workbooks
    uv run --extra excel --extra browser python scripts/compare_excel_backends.py my.xlsx    # also your own (copies)

On a PC with Excel (Windows or macOS), both backends run; elsewhere only openpyxl (the report says
so). Your own workbooks are copied first: the originals are never touched. The report is printed
and written to excel-backend-comparison/report.md, next to every result (edited workbooks, the
pictures view_excel produced) so you can open them and look.

What is compared, for each backend:
  1. Editing a workbook that holds a chart, a picture, a pivot table, a conditional format, a
     drop-down list, a comment and a formula: which of them are still in the file afterwards, and
     whether the formula's result is up to date.
  2. Text that looks like a number ("00123") stays text.
  3. Creating a clean table (the schools list), formatting it (header, widths, wrap, frozen header,
     filters) and viewing it.
  4. Viewing a sheet with a chart.
  5. add_chart, add_table, add_pivot_table: a chart, an Excel table and a pivot table (Excel only) added to a list of expenses,
     then the file checked and the sheets viewed.
  6. How long each step takes.
"""

from __future__ import annotations

import argparse
import shutil
import struct
import sys
import tempfile
import time
import traceback
import zipfile
import zlib
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402
from openpyxl.chart import BarChart, Reference  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.formatting.rule import CellIsRule  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from coding_agent import backups, state  # noqa: E402
from coding_agent.tools import documents, excel_look, excel_xl  # noqa: E402
from coding_agent.tools.excel_add import tool_add_chart, tool_add_pivot_table, tool_add_table  # noqa: E402
from coding_agent.ui import UI  # noqa: E402

OUT = Path("excel-backend-comparison")
PARTS = {  # what an .xlsx holds, by the folders of its zip
    "charts": "xl/charts/chart", "pictures": "xl/media/", "drawings": "xl/drawings/drawing",
    "pivot tables": "xl/pivotTables/", "pivot caches": "xl/pivotCache/", "slicers": "xl/slicers/",
    "comments": "xl/comments", "threaded comments": "xl/threadedComments/", "form controls": "xl/ctrlProps/",
    "macros": "xl/vbaProject.bin", "Excel tables": "xl/tables/",
}
SCHOOLS = [  # the list from the conversation, split into one piece of information per column
    ("Tunisie", "Faculté Privée de Montplaisir", "Tunis", "+216 71 902 323; +216 71 901 660", "contact@umt.ens.tn", "", "Trouvé"),
    ("Tunisie", "Institut Privé des Hautes Etudes de Tunis (IPHET)", "Tunis", "+216 71 841 855", "iehet@gnet.tn", "", "Trouvé"),
    ("Algérie", "MDI-Algiers Business School", "Alger", "", "", "mdi-alger.dz", "À vérifier"),
    ("Sénégal", "Institut Supérieur de Technologie Industrielle (ISTI)", "Dakar", "+221 33 824 38 39", "", "",
     "Indicatif régional supposé"),
    ("Maroc", "Université Internationale de Rabat (UIR)", "Rabat", "", "", "uir.ac.ma", "À vérifier"),
    ("Côte d'Ivoire", "Université Adama Sanogo (UAS)", "Abidjan", "", "", "", "À vérifier"),
]


class YesUI(UI):
    """Approves everything (the files are copies)."""

    def confirm(self, question, choices=("yes", "no")):
        return "yes"

    def ask_text(self, prompt, multiline=False):
        return ""


def png(width=60, height=30, rgb=(46, 117, 182)) -> bytes:
    """A small plain PNG, without an imaging library."""
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))

    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def sample(path: Path, with_excel: bool) -> list[str]:
    """A workbook holding what openpyxl is known to damage. Returns what could be put in it."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Pays", "Ecoles", "Budget"])
    for row in [["Tunisie", 14, 1200], ["Sénégal", 22, 800], ["Maroc", 6, 1500], ["Algérie", 4, 600]]:
        ws.append(row)
    ws["A7"], ws["B7"] = "Total", "=SUM(B2:B5)"
    ws.conditional_formatting.add("C2:C5", CellIsRule(operator="greaterThan", formula=["1000"],
                                                      fill=PatternFill("solid", fgColor="FFC7CE")))
    dv = DataValidation(type="list", formula1='"Trouvé,À vérifier"')
    ws.add_data_validation(dv)
    dv.add("D2:D5")
    ws["A1"].comment = Comment("Country of the school", "agent test")
    chart = BarChart()
    chart.title = "Ecoles par pays"
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
    ws.add_chart(chart, "F2")
    wb.save(path)
    made = ["a chart (openpyxl)", "a conditional format", "a drop-down list", "a comment", "a formula"]
    if with_excel:  # what only Excel can add: a picture, a pivot table and a chart styled by Excel
        try:
            made += excel_extras(path)
        except Exception as e:
            made.append(f"(Excel could not add its extras: {type(e).__name__}: {e})")
    return made


def excel_extras(path: Path) -> list[str]:
    app = excel_xl.app()
    book = app.books.open(str(path))
    made = []
    try:
        ws = book.sheets["Data"]
        with tempfile.TemporaryDirectory() as tmp:
            picture = Path(tmp) / "logo.png"
            picture.write_bytes(png())
            ws.pictures.add(str(picture), left=ws.range("F20").left, top=ws.range("F20").top)
            made.append("a picture (Excel)")
        chart = ws.charts.add(left=ws.range("M2").left, top=ws.range("M2").top)
        chart.set_source_data(ws.range("A1:C5"))
        chart.chart_type = "column_clustered"
        made.append("a chart made by Excel")
        try:  # a pivot table, through Excel's own API (Windows)
            pivot_sheet = book.sheets.add("Pivot")
            cache = book.api.PivotCaches().Create(1, ws.range("A1:C5").api)
            table = cache.CreatePivotTable(pivot_sheet.range("A3").api, "SchoolsPivot")
            table.AddDataField(table.PivotFields("Ecoles"))
            made.append("a pivot table (Excel)")
        except Exception as e:
            made.append(f"(no pivot table: {type(e).__name__})")
        book.save()
    finally:
        book.close()
    return made


def inventory(path: Path) -> dict:
    """What the file holds: its parts, the chart files' contents, and a few sheet features."""
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        charts = {n: z.read(n) for n in names if n.startswith("xl/charts/chart")}
    found = {label: sum(n.startswith(prefix) for n in names) for label, prefix in PARTS.items()}
    wb = openpyxl.load_workbook(path)
    ws = wb["Data"]
    found["conditional formats"] = sum(len(r.rules) for r in ws.conditional_formatting)
    found["drop-down lists"] = len(ws.data_validations.dataValidation)
    found["comments in cells"] = sum(1 for row in ws.iter_rows() for c in row if c.comment)
    found["_charts"] = charts
    total = openpyxl.load_workbook(path, data_only=True)["Data"]["B7"].value
    found["formula result (B7)"] = total
    return found


def timed(results: dict, label: str, action):
    start = time.perf_counter()
    try:
        value = action()
        results[label] = f"{time.perf_counter() - start:.1f} s"
        return value
    except Exception as e:
        results[label] = f"failed: {type(e).__name__}: {str(e)[:150]}"
        traceback.print_exc()
        return None


def run(backend: str, work: Path, own: list[Path], with_excel: bool) -> dict:
    folder = OUT / backend
    shutil.rmtree(folder, ignore_errors=True)
    folder.mkdir(parents=True)
    documents.EXCEL_BACKEND = backend
    state.set_workspace(folder.resolve(), f"compare-{backend}", work / "memory")
    state.auto_mode = True
    state.ui = YesUI()
    r: dict = {}

    # 1. edit a workbook holding everything
    shutil.copy2(work / "sample.xlsx", folder / "sample.xlsx")
    before = inventory(folder / "sample.xlsx")
    timed(r, "edit: time", lambda: documents.tool_edit_excel(
        "sample.xlsx", [{"sheet": "Data", "cell": "B2", "value": 15}, {"sheet": "Data", "cell": "A9", "value": "00123"}]))
    after = inventory(folder / "sample.xlsx")
    for key in list(PARTS) + ["conditional formats", "drop-down lists", "comments in cells"]:
        if before[key]:
            r[f"edit: {key} kept"] = f"{after[key]}/{before[key]}" + (" ✅" if after[key] >= before[key] else " ❌")
    same = [n for n, data in before["_charts"].items() if after["_charts"].get(n) == data]
    if before["_charts"]:
        r["edit: charts unchanged (byte for byte)"] = (f"{len(same)}/{len(before['_charts'])}"
                                                       + (" ✅" if len(same) == len(before["_charts"]) else " ⚠️ rewritten"))
    total = after["formula result (B7)"]
    r["edit: formula result updated (B7 = 47)"] = ("47 ✅" if total == 47 else
                                                  f"{total!r} ❌ (stale until opened in Excel)")
    text = openpyxl.load_workbook(folder / "sample.xlsx")["Data"]["A9"].value
    r["edit: '00123' kept as text"] = "✅" if text == "00123" else f"❌ became {text!r}"

    # 2. a clean table: create, format, view
    rows = [["Pays", "Établissement", "Ville", "Téléphone", "Email", "Site web", "Statut"], *SCHOOLS]
    changes = [{"sheet": "Écoles", "cell": f"{'ABCDEFG'[c]}{r_ + 1}", "value": v}
               for r_, row in enumerate(rows) for c, v in enumerate(row) if v != ""]
    timed(r, "table: create", lambda: documents.tool_edit_excel("schools.xlsx", changes, create_sheets=["Écoles"]))
    timed(r, "table: format", lambda: [
        excel_look.tool_format_excel("schools.xlsx", "A1:G1", sheet="Écoles", bold=True, fill="#DDEBF7",
                                     freeze="A2", autofilter=True),
        excel_look.tool_format_excel("schools.xlsx", "A:G", sheet="Écoles", column_width="auto"),
        excel_look.tool_format_excel("schools.xlsx", "B:B", sheet="Écoles", column_width=45, wrap=True)])
    view = timed(r, "table: view", lambda: excel_look.tool_view_excel("schools.xlsx", sheet="Écoles"))
    r["table: view gives"] = save_view(view, folder / "schools-view")

    # 3. view a sheet with charts
    chart_view = timed(r, "chart: view", lambda: excel_look.tool_view_excel("sample.xlsx", sheet="Data"))
    r["chart: view gives"] = save_view(chart_view, folder / "sample-view")

    # 4. add_chart, add_table, add_pivot_table: a chart, an Excel table and a pivot table on a list of expenses
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Expenses"
    ws.append(["Date", "Supplier", "Month", "Amount"])
    for row in [["2026-01-03", "Acme", "Jan", 120], ["2026-01-15", "Beta", "Jan", 80],
                ["2026-02-02", "Acme", "Feb", 200], ["2026-02-20", "Gamma", "Feb", 45]]:
        ws.append(row)
    wb.save(folder / "expenses.xlsx")
    timed(r, "add: table", lambda: tool_add_table("expenses.xlsx", "A1:D5", name="Expenses"))
    timed(r, "add: chart", lambda: tool_add_chart("expenses.xlsx", "B1:D5", chart_type="column",
                                                  title="Spend by supplier", anchor="F2"))
    pivot = timed(r, "add: pivot table", lambda: tool_add_pivot_table(
        "expenses.xlsx", "A1:D5", rows=["Supplier"], columns=["Month"],
        values=[{"field": "Amount", "summary": "sum"}]))
    if pivot is None and backend == "openpyxl":
        r["add: pivot table"] = "refused (needs Excel) ✅"
    parts = inventory_parts(folder / "expenses.xlsx")
    r["add: in the file"] = ", ".join(f"{k} {parts[k]}" for k in ("Excel tables", "charts", "pivot tables"))
    added = timed(r, "add: view", lambda: excel_look.tool_view_excel("expenses.xlsx", sheet="Expenses"))
    r["add: view gives"] = save_view(added, folder / "expenses-view")
    if backend == "xlwings":
        pivot_view = timed(r, "add: pivot view", lambda: excel_look.tool_view_excel("expenses.xlsx", sheet="Pivot"))
        r["add: pivot view gives"] = save_view(pivot_view, folder / "pivot-view")

    # 5. your own workbooks: one value written below the data of the first sheet
    for path in own:
        copy = folder / f"own-{path.name}"
        shutil.copy2(path, copy)
        first = openpyxl.load_workbook(copy, read_only=True).sheetnames[0]
        rows_used = openpyxl.load_workbook(copy)[first].max_row
        kinds = {k: v for k, v in inventory_parts(copy).items() if v}
        timed(r, f"{path.name}: edit", lambda: documents.tool_edit_excel(
            copy.name, [{"sheet": first, "cell": f"A{rows_used + 2}", "value": "agent test"}]))
        now = inventory_parts(copy)
        r[f"{path.name}: kept"] = ", ".join(f"{k} {now.get(k, 0)}/{v}" for k, v in kinds.items()) or "(nothing special in it)"
    return r


def inventory_parts(path: Path) -> dict:
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
    return {label: sum(n.startswith(prefix) for n in names) for label, prefix in PARTS.items()}


def save_view(blocks, stem: Path) -> str:
    if not blocks:
        return "nothing"
    import base64
    block = blocks[1]
    data = base64.b64decode(block["source"]["data"])
    suffix = ".pdf" if block["type"] == "document" else ".png"
    stem.with_suffix(suffix).write_bytes(data)
    return f"{'the PDF Excel prints' if suffix == '.pdf' else 'a picture'} ({len(data) // 1024} KB): {stem.with_suffix(suffix)}"


def report(results: dict[str, dict], made: list[str], notes: list[str]) -> str:
    backends = list(results)
    keys = [k for k in dict.fromkeys(k for r in results.values() for k in r)]
    lines = ["# Excel backends: openpyxl vs xlwings", "", "Sample workbook: " + ", ".join(made) + ".", ""]
    lines += [f"- {n}" for n in notes] + ([""] if notes else [])
    lines.append("| Check | " + " | ".join(backends) + " |")
    lines.append("|---|" + "---|" * len(backends))
    for key in keys:
        lines.append(f"| {key} | " + " | ".join(str(results[b].get(key, "")) for b in backends) + " |")
    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("workbooks", nargs="*", type=Path, help="Your own .xlsx/.xlsm files to test (copied first).")
    args = parser.parse_args()
    OUT.mkdir(exist_ok=True)
    notes = []
    with_excel = excel_xl.installed() and excel_xl.available()
    if not with_excel:
        notes.append("Excel could not be driven here (xlwings not installed, or no Excel): only openpyxl was run. "
                     "Run this on a PC with Excel: uv run --extra excel --extra browser python scripts/compare_excel_backends.py")
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        backups.BACKUP_HOME = work / "backups"
        made = sample(work / "sample.xlsx", with_excel)
        results = {"openpyxl": run("openpyxl", work, args.workbooks, with_excel)}
        if with_excel:
            results["xlwings"] = run("xlwings", work, args.workbooks, with_excel)
    excel_xl.quit_excel()
    text = report(results, made, notes)
    (OUT / "report.md").write_text(text, encoding="utf-8")
    print(text)
    print(f"Results (workbooks, views): {OUT.resolve()}")


if __name__ == "__main__":
    main()
