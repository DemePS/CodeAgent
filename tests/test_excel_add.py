"""add_chart, add_table, add_pivot_table: charts, Excel tables and pivot tables, with openpyxl and with Excel (played by fake_xlwings)."""

import sys

import openpyxl
import pytest

from coding_agent import backups, state
from coding_agent.common import ToolError
from coding_agent.tools import documents, excel_xl
from coding_agent.tools.excel_add import tool_add_chart, tool_add_pivot_table, tool_add_table

import fake_xlwings


@pytest.fixture
def backup_home(tmp_path, monkeypatch):
    monkeypatch.setattr(backups, "BACKUP_HOME", tmp_path / "backups")


@pytest.fixture
def xl(monkeypatch, backup_home):
    """The xlwings backend, with Excel played by fake_xlwings."""
    monkeypatch.setitem(sys.modules, "xlwings", fake_xlwings)
    monkeypatch.setattr(fake_xlwings, "apps", [])
    monkeypatch.setattr(excel_xl, "_app", None)
    monkeypatch.setattr(excel_xl, "_available", None)
    monkeypatch.setattr(documents, "EXCEL_BACKEND", "xlwings")
    return fake_xlwings


@pytest.fixture
def openpyxl_backend(monkeypatch, backup_home):
    monkeypatch.setattr(documents, "EXCEL_BACKEND", "openpyxl")


@pytest.fixture
def expenses(workspace):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in [["Date", "Supplier", "Month", "Amount"], ["2026-01-03", "Acme", "Jan", 120],
                ["2026-01-15", "Beta", "Jan", 80], ["2026-02-02", "Acme", "Feb", 200]]:
        ws.append(row)
    wb.save(workspace / "expenses.xlsx")
    return workspace / "expenses.xlsx"


# --- charts

@pytest.mark.parametrize("chart_type, cls", [("column", "BarChart"), ("bar", "BarChart"), ("line", "LineChart"),
                                             ("area", "AreaChart"), ("scatter", "ScatterChart")])
def test_chart_with_openpyxl(workspace, expenses, ui, openpyxl_backend, chart_type, cls):
    ui.answers = ["yes"]
    result = tool_add_chart("expenses.xlsx", "C1:D4", chart_type=chart_type, title="Spend")
    ws = openpyxl.load_workbook(expenses)["Data"]
    assert [type(c).__name__ for c in ws._charts] == [cls]
    assert "Spend" in result and backups.versions("expenses.xlsx")
    listed = documents.tool_read_excel("expenses.xlsx")
    assert "chart 'Spend' at F1" in listed and "'Data'!$D$2:$D$4" in listed  # beside the data, plotting Amount


def test_pie_chart_takes_two_columns(workspace, expenses, ui, openpyxl_backend):
    with pytest.raises(ToolError, match="one series"):
        tool_add_chart("expenses.xlsx", "B1:D4", chart_type="pie")
    ui.answers = ["yes"]
    tool_add_chart("expenses.xlsx", "C1:D4", chart_type="pie", anchor="H10")
    ws = openpyxl.load_workbook(expenses)["Data"]
    assert type(ws._charts[0]).__name__ == "PieChart"
    assert "at H10" in documents.tool_read_excel("expenses.xlsx")


def test_chart_with_excel(workspace, expenses, ui, xl, monkeypatch):
    monkeypatch.setattr(state, "auto_mode", True)
    result = tool_add_chart("expenses.xlsx", "C1:D4", chart_type="column", title="Spend", anchor="F2")
    chart = excel_xl._app.charts[0]
    assert chart.source == "C1:D4" and chart.chart_type == "column_clustered"
    assert chart.position == {"left": 5 * 48.0, "top": 15.0, "width": 480, "height": 288}
    assert chart.api[1].ChartTitle.Text == "Spend" and chart.api[1].PlotBy == 2
    assert "Not applied" not in result and "may have been lost" not in result
    assert excel_xl._app.saved == ["expenses.xlsx"]


# --- tables

def test_table_with_openpyxl(workspace, expenses, ui, openpyxl_backend):
    ui.answers = ["yes"]
    tool_add_table("expenses.xlsx", "A1:D4", name="Expenses")
    ws = openpyxl.load_workbook(expenses)["Data"]
    assert ws.tables["Expenses"].ref == "A1:D4"
    assert [c.name for c in ws.tables["Expenses"].tableColumns] == ["Date", "Supplier", "Month", "Amount"]
    ui.answers = ["yes"]
    with pytest.raises(ToolError, match="overlaps the Excel table 'Expenses'"):
        tool_add_table("expenses.xlsx", "B2:C3")


def test_table_names_and_headers(workspace, expenses, ui, openpyxl_backend):
    for name in ("my table", "A1", "R1C1", "1st"):
        with pytest.raises(ToolError, match="name"):
            tool_add_table("expenses.xlsx", "A1:D4", name=name)
    wb = openpyxl.load_workbook(expenses)
    wb["Data"]["C1"] = "supplier"  # same header as B1, but for case
    wb.save(expenses)
    with pytest.raises(ToolError, match="must differ"):
        tool_add_table("expenses.xlsx", "A1:D4")
    ui.answers = ["yes"]
    tool_add_table("expenses.xlsx", "A1:B4")  # no name given: Table1
    assert "Table1" in openpyxl.load_workbook(expenses)["Data"].tables


def test_table_with_excel(workspace, expenses, ui, xl, monkeypatch):
    monkeypatch.setattr(state, "auto_mode", True)
    tool_add_table("expenses.xlsx", "A1:D4", name="Expenses")
    table = openpyxl.load_workbook(expenses)["Data"].tables["Expenses"]
    assert table.ref == "A1:D4" and table.tableStyleInfo.name == "TableStyleMedium2"


# --- pivot tables

def test_pivot_table_with_excel(workspace, expenses, ui, xl):
    ui.answers = ["yes"]
    result = tool_add_pivot_table("expenses.xlsx", "A1:D4", rows=["supplier"], columns=["Month"],
                            values=[{"field": "Amount", "summary": "sum"}, {"field": "Amount", "summary": "count"}])
    panel = ui.of("panel")[-1]
    assert panel[1] == "Add a pivot table to expenses.xlsx"
    assert panel[2] == ["pivot table 'Pivot1' from Data!A1:D4, at new sheet Pivot!A3", "rows: Supplier",
                        "columns: Month", "values: sum of Amount, count of Amount"]
    pivot = excel_xl._app.pivots[0]
    assert pivot.name == "Pivot1" and pivot.destination.Address == "A3"
    assert (pivot.fields["Supplier"].Orientation, pivot.fields["Month"].Orientation) == (1, 2)
    assert pivot.data == [("Amount", "Sum of Amount", -4157), ("Amount", "Count of Amount", -4112)]
    assert openpyxl.load_workbook(expenses).sheetnames == ["Data", "Pivot"]
    assert "Previous version saved" in result


def test_pivot_table_on_an_existing_sheet(workspace, expenses, ui, xl):
    wb = openpyxl.load_workbook(expenses)
    wb.create_sheet("Summary")["B2"] = "taken"
    wb.save(expenses)
    with pytest.raises(ToolError, match="give anchor"):
        tool_add_pivot_table("expenses.xlsx", "A1:D4", sheet="Data", rows=["Supplier"],
                       values=["Amount"], target_sheet="Summary")
    with pytest.raises(ToolError, match="not empty"):
        tool_add_pivot_table("expenses.xlsx", "A1:D4", sheet="Data", rows=["Supplier"],
                       values=["Amount"], target_sheet="Summary", anchor="B2")
    ui.answers = ["yes"]
    tool_add_pivot_table("expenses.xlsx", "A1:D4", sheet="Data", rows=["Supplier"], values=["Amount"],
                   target_sheet="Summary", anchor="D5", name="BySupplier")
    pivot = excel_xl._app.pivots[0]
    assert pivot.name == "BySupplier" and pivot.destination.Address == "D5"
    assert openpyxl.load_workbook(expenses).sheetnames == ["Data", "Summary"]


def test_pivot_table_checks(workspace, expenses, ui, xl):
    for kwargs, refusal in [({"rows": [], "values": ["Amount"]}, "rows: the header"),
                            ({"rows": ["Vendor"], "values": ["Amount"]}, "'Vendor' is not a header"),
                            ({"rows": ["Supplier"], "values": []}, "values: what to total"),
                            ({"rows": ["Supplier"], "columns": ["Supplier"], "values": ["Amount"]}, "not both"),
                            ({"rows": ["Supplier"], "values": [{"field": "Amount", "summary": "median"}]}, "summary")]:
        with pytest.raises(ToolError, match=refusal):
            tool_add_pivot_table("expenses.xlsx", "A1:D4", **kwargs)
    assert excel_xl._app is None or not excel_xl._app.pivots


def test_pivot_table_needs_excel(workspace, expenses, ui, openpyxl_backend):
    with pytest.raises(ToolError, match="Pivot tables need Excel"):
        tool_add_pivot_table("expenses.xlsx", "A1:D4", rows=["Supplier"], values=["Amount"])
    assert not backups.versions("expenses.xlsx")


# --- what the three tools share

def test_rejected_or_disallowed_changes_nothing(workspace, expenses, ui, openpyxl_backend, monkeypatch):
    before = expenses.read_bytes()
    ui.answers = ["no", "not now"]
    with pytest.raises(ToolError, match="rejected.*User feedback: not now"):
        tool_add_table("expenses.xlsx", "A1:D4")
    assert expenses.read_bytes() == before and not backups.versions("expenses.xlsx")
    monkeypatch.setattr(state, "excel_allow_format", False)
    with pytest.raises(ToolError, match="not allowed in this application"):
        tool_add_chart("expenses.xlsx", "C1:D4", chart_type="line")
    monkeypatch.setattr(state, "excel_allow_format", True)
    monkeypatch.setattr(state, "excel_edit_sheets", {expenses.resolve(): {"Other"}})
    with pytest.raises(ToolError, match="Only these sheets"):
        tool_add_table("expenses.xlsx", "A1:D4")


def test_input_checks(workspace, expenses, ui, openpyxl_backend):
    for tool, kwargs, refusal in [(tool_add_table, {"source": "A:D"}, "header row"),
                                  (tool_add_table, {"source": "A1:D1"}, "at least one row"),
                                  (tool_add_chart, {"source": "C1:D4", "chart_type": "donut"}, "chart_type"),
                                  (tool_add_chart, {"source": "D1:D4", "chart_type": "line"}, "at least two columns"),
                                  (tool_add_chart, {"source": "C1:D4", "chart_type": "line", "anchor": "here"}, "anchor"),
                                  (tool_add_table, {"source": "A1:D4", "sheet": "Nope"}, "No sheet")]:
        with pytest.raises(ToolError, match=refusal):
            tool("expenses.xlsx", **kwargs)


def test_the_three_tools_are_offered_with_their_own_fields():
    from coding_agent.schemas import TOOLS
    from coding_agent.tools import TOOL_HANDLERS

    schemas = {t["name"]: t for t in TOOLS}
    fields = {name: set(schemas[name]["input_schema"]["properties"]) for name in ("add_chart", "add_table", "add_pivot_table")}
    assert fields["add_table"] == {"path", "source", "sheet", "name"}
    assert "rows" not in fields["add_chart"] and "chart_type" not in fields["add_pivot_table"]
    assert all(name in TOOL_HANDLERS for name in fields) and "excel_add" not in schemas


def test_openpyxl_save_warns_when_the_workbook_has_charts(workspace, expenses, ui, openpyxl_backend, monkeypatch):
    ui.answers = ["yes"]
    tool_add_chart("expenses.xlsx", "C1:D4", chart_type="line")
    monkeypatch.setattr(state, "auto_mode", True)  # a lossy save still asks
    ui.answers = ["yes"]
    result = tool_add_table("expenses.xlsx", "A1:D4")
    assert any("charts" in w[1] for w in ui.of("warning")) and ui.of("confirm")
    assert "may have been lost" in result
