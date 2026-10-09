"""The two Excel backends: openpyxl (rewrites the file) and xlwings (Excel does the work), and the
tools on top of them: read_excel's charts and long cells, edit_excel, view_excel, format_excel."""

import sys

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference

from coding_agent import backups, state
from coding_agent.common import ToolError
from coding_agent.tools import documents, excel_look, excel_xl

import fake_xlwings


@pytest.fixture
def backup_home(tmp_path, monkeypatch):
    monkeypatch.setattr(backups, "BACKUP_HOME", tmp_path / "backups")


@pytest.fixture
def xl(monkeypatch, backup_home):
    """The xlwings backend, with Excel played by fake_xlwings."""
    monkeypatch.setitem(sys.modules, "xlwings", fake_xlwings)
    monkeypatch.setattr(fake_xlwings, "apps", [])
    monkeypatch.setattr(excel_xl, "_app", None)
    monkeypatch.setattr(excel_xl, "_available", None)
    monkeypatch.setattr(documents, "EXCEL_BACKEND", "xlwings")
    return fake_xlwings


@pytest.fixture
def openpyxl_backend(monkeypatch, backup_home):
    monkeypatch.setattr(documents, "EXCEL_BACKEND", "openpyxl")


@pytest.fixture
def charted(workspace):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Schools"
    for row in [["Pays", "Ecoles"], ["Tunisie", 14], ["Senegal", 22], ["Maroc", 6]]:
        ws.append(row)
    ws["A6"] = "Guidance: " + "x" * 300 + " END"
    chart = BarChart()
    chart.title = "Ecoles par pays"
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    ws.add_chart(chart, "D2")
    wb.save(workspace / "schools.xlsx")
    return workspace / "schools.xlsx"


# --- choosing the backend

def test_auto_uses_excel_only_when_it_can_be_driven(monkeypatch, workspace):
    monkeypatch.setattr(documents, "EXCEL_BACKEND", "auto")
    monkeypatch.setattr(excel_xl, "_available", None)
    monkeypatch.setitem(sys.modules, "xlwings", None)  # not installed
    assert documents.excel_backend() == "openpyxl"
    monkeypatch.setitem(sys.modules, "xlwings", fake_xlwings)
    monkeypatch.setattr(excel_xl, "_available", None)
    monkeypatch.setattr(excel_xl, "_app", None)
    assert documents.excel_backend() == "xlwings"
    monkeypatch.setattr(documents, "EXCEL_BACKEND", "openpyxl")
    assert documents.excel_backend() == "openpyxl"


def test_asking_for_xlwings_without_it_is_explained(monkeypatch):
    monkeypatch.setattr(documents, "EXCEL_BACKEND", "xlwings")
    monkeypatch.setitem(sys.modules, "xlwings", None)
    with pytest.raises(ToolError, match="xlwings is not installed"):
        documents.excel_backend()


# --- read_excel: charts and long cells (both backends read with openpyxl)

def test_read_lists_the_charts_and_long_cells_can_be_read_whole(charted):
    text = documents.tool_read_excel("schools.xlsx")
    assert "chart: bar chart 'Ecoles par pays' at D2, plotting 'Schools'!$A$2:$A$4, 'Schools'!$B$2:$B$4" in text
    assert "A6=Guidance: " in text and "more characters: read this cell alone" in text and "END" not in text
    whole = documents.tool_read_excel("schools.xlsx", range="A6")
    assert whole.rstrip().endswith("x END")


# --- edit_excel

def test_excel_saves_the_changes_and_keeps_what_openpyxl_would_lose(workspace, charted, ui, xl):
    ui.answers = ["yes"]
    result = documents.tool_edit_excel("schools.xlsx", [
        {"cell": "B5", "value": "=SUM(B2:B4)"},
        {"cell": "A7", "value": "00123"},  # text that Excel would turn into a number
        {"cell": "A8", "value": 42},
    ])
    assert not ui.of("warning")  # the chart is kept: no loss warning with Excel
    assert "Excel recalculated the formulas" in result and "may have been lost" not in result
    ws = openpyxl.load_workbook(charted)["Schools"]
    assert ws["B5"].value == "=SUM(B2:B4)" and ws["A7"].value == "00123" and ws["A8"].value == 42
    [app] = [excel_xl._app]
    assert app.saved == ["schools.xlsx"] and app.open_books == []  # saved by Excel, then closed
    assert app.visible is False and app.display_alerts is False
    assert backups.versions("schools.xlsx")  # the previous version is kept


def test_openpyxl_warns_that_the_chart_may_be_lost(workspace, charted, ui, openpyxl_backend):
    ui.answers = ["yes"]
    result = documents.tool_edit_excel("schools.xlsx", [{"cell": "A8", "value": 42}])
    assert any("charts" in w[1] for w in ui.of("warning")) and "may have been lost" in result


def test_the_checks_are_the_same_with_excel(workspace, ui, xl, monkeypatch):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q"
    ws.append(["#", "Question", "Answer"])
    ws["C3"] = "=B2"
    wb.create_sheet("Instructions")
    wb.save(workspace / "q.xlsx")
    p = (workspace / "q.xlsx").resolve()
    monkeypatch.setattr(state, "excel_protect_formulas", True)
    monkeypatch.setattr(state, "excel_max_columns", {p: {"Q": 3}})
    monkeypatch.setattr(state, "excel_edit_sheets", {p: {"Q"}})
    for changes, refusal in [
        ([{"sheet": "Q", "cell": "C2", "value": "=MID(B2,1,2)"}], "formulas may not be written"),
        ([{"sheet": "Q", "cell": "C3", "value": None}], "holds a formula"),
        ([{"sheet": "Q", "cell": "I2", "value": "x"}], "outside the sheet's columns"),
        ([{"sheet": "Instructions", "cell": "B2", "value": "x"}], "Only these sheets"),
        ([{"cell": "C2", "value": "x"}], "give 'sheet' for every change"),
    ]:
        with pytest.raises(ToolError, match=refusal):
            documents.tool_edit_excel("q.xlsx", changes)
    assert excel_xl._app is None or excel_xl._app.saved == []  # Excel never touched the file
    ui.answers = ["no", ""]
    with pytest.raises(ToolError, match="rejected"):
        documents.tool_edit_excel("q.xlsx", [{"sheet": "Q", "cell": "C2", "value": "Yes"}])
    assert openpyxl.load_workbook(p)["Q"]["C2"].value is None


def test_excel_creates_a_workbook_with_the_same_sheets(workspace, ui, xl):
    ui.answers = ["yes"]
    documents.tool_edit_excel("new.xlsx", [{"sheet": "Data", "cell": "A1", "value": "Pays"}], create_sheets=["Data"])
    wb = openpyxl.load_workbook(workspace / "new.xlsx")
    assert wb.sheetnames == ["Sheet", "Data"] and wb["Data"]["A1"].value == "Pays"  # not "Feuil1"


def test_a_workbook_open_in_the_persons_excel(workspace, charted, ui, xl, monkeypatch):
    theirs = xl.App()
    book = theirs.books.open(str(charted))
    monkeypatch.setattr(xl, "apps", [theirs])
    monkeypatch.setattr(state, "auto_mode", True)
    book.api.Saved = False
    with pytest.raises(ToolError, match="open in Excel with unsaved changes"):
        documents.tool_edit_excel("schools.xlsx", [{"cell": "A8", "value": 1}])
    book.api.Saved = True
    documents.tool_edit_excel("schools.xlsx", [{"cell": "A8", "value": 1}])
    assert theirs.saved == ["schools.xlsx"] and book in theirs.open_books  # written in their window, left open


# --- format_excel

def test_format_with_openpyxl(workspace, charted, ui, openpyxl_backend):
    ui.answers = ["yes"]
    result = excel_look.tool_format_excel("schools.xlsx", "A1:B1", bold=True, fill="#DDEBF7", freeze="A2",
                                          autofilter=True)
    ui.answers = ["yes"]
    excel_look.tool_format_excel("schools.xlsx", "A:B", column_width="auto", wrap=True)
    ws = openpyxl.load_workbook(charted)["Schools"]
    assert ws["A1"].font.b and ws["B1"].fill.fgColor.rgb == "FFDDEBF7" and ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:B1" and ws["A6"].alignment.wrap_text
    assert ws.column_dimensions["A"].width == 80 and 6 <= ws.column_dimensions["B"].width < 12
    assert "may have been lost" in result and backups.versions("schools.xlsx")


def test_format_with_excel(workspace, charted, ui, xl, monkeypatch):
    monkeypatch.setattr(state, "auto_mode", True)
    result = excel_look.tool_format_excel("schools.xlsx", "A1:B1", bold=True, fill="#DDEBF7", freeze="A2",
                                          autofilter=True, column_width="auto", wrap=True, border=True)
    assert "may have been lost" not in result and "Not applied" not in result
    ws = openpyxl.load_workbook(charted)["Schools"]
    assert ws["A1"].font.b and ws["B1"].fill.fgColor.rgb == "FFDDEBF7"
    assert excel_xl._app.api.ActiveWindow.FreezePanes is True and excel_xl._app.api.ActiveWindow.SplitRow == 1


def test_format_checks(workspace, charted, ui, openpyxl_backend, monkeypatch):
    for kwargs, refusal in [({"fill": "blue"}, "#RRGGBB"), ({"column_width": "wide"}, "column_width"),
                            ({"freeze": "row 2"}, "freeze"), ({}, "No formatting given")]:
        with pytest.raises(ToolError, match=refusal):
            excel_look.tool_format_excel("schools.xlsx", "A1:B1", **kwargs)
    with pytest.raises(ToolError, match="Invalid range"):
        excel_look.tool_format_excel("schools.xlsx", "header", bold=True)
    monkeypatch.setattr(state, "excel_allow_format", False)  # an application filling existing workbooks
    with pytest.raises(ToolError, match="not allowed in this application"):
        excel_look.tool_format_excel("schools.xlsx", "A1:B1", bold=True)


# --- view_excel

def test_view_with_excel_gives_the_printed_sheet_or_a_picture(workspace, charted, xl, monkeypatch):
    monkeypatch.setattr(excel_look, "uses_deepseek", lambda: False)
    sheet = excel_look.tool_view_excel("schools.xlsx")
    assert "as Excel shows it" in sheet[0]["text"] and sheet[1]["type"] == "document"
    picture = excel_look.tool_view_excel("schools.xlsx", range="a1:b4")
    assert picture[1]["type"] == "image" and picture[1]["source"]["media_type"] == "image/png"


def test_view_under_deepseek_renders_the_printed_pages_to_pictures(workspace, charted, xl, monkeypatch):
    monkeypatch.setattr(excel_look, "uses_deepseek", lambda: True)
    monkeypatch.setattr(excel_look, "render_pdf_pages", lambda p, pages: [(b"\x89PNG fake", "image/png") for _ in pages])
    sheet = excel_look.tool_view_excel("schools.xlsx")
    assert "as Excel shows it" in sheet[0]["text"]
    assert "picture of the printed sheet" in sheet[1]["text"]
    assert sheet[2]["type"] == "image" and sheet[2]["source"]["media_type"] == "image/png"


def test_view_under_deepseek_without_the_renderer_says_how_to_install_it(workspace, charted, xl, monkeypatch):
    def missing(p, pages):
        raise ToolError('needs the optional renderer: pip install "codeagent[pdf-image]"')

    monkeypatch.setattr(excel_look, "uses_deepseek", lambda: True)
    monkeypatch.setattr(excel_look, "render_pdf_pages", missing)
    with pytest.raises(ToolError, match=r"codeagent\[pdf-image\]"):
        excel_look.tool_view_excel("schools.xlsx")


def test_view_without_excel_draws_the_cells(charted, openpyxl_backend, monkeypatch):
    wb = openpyxl.load_workbook(charted)
    ws = wb["Schools"]
    ws["A1"].font = openpyxl.styles.Font(b=True)
    ws.merge_cells("A9:C9")
    page = excel_look.sheet_html(ws, None)
    assert "font-weight:bold" in page and 'colspan="3"' in page and ">Tunisie<" in page
    monkeypatch.setattr(excel_look, "screenshot", lambda html: b"\x89PNG fake")
    view = excel_look.tool_view_excel("schools.xlsx")
    assert "drawn without Excel" in view[0]["text"] and "not drawn: chart: bar chart 'Ecoles par pays'" in view[0]["text"]
    assert view[1]["type"] == "image"


def test_freezing_resets_a_stale_scroll_position(workspace, ui, openpyxl_backend, monkeypatch):
    import re
    import zipfile

    wb = openpyxl.Workbook()
    for i in range(100):
        wb.active.append([i])
    wb.active.sheet_view.topLeftCell = "A43"  # left scrolled down by the person
    wb.save(workspace / "long.xlsx")
    monkeypatch.setattr(state, "auto_mode", True)
    excel_look.tool_format_excel("long.xlsx", "A1:A1", freeze="A2")
    xml = zipfile.ZipFile(workspace / "long.xlsx").read("xl/worksheets/sheet1.xml").decode()
    view = re.search(r"<sheetView .*?</sheetView>", xml).group()
    assert 'topLeftCell="A1"' in view and '<pane ySplit="1" topLeftCell="A2"' in view and 'activeCell="A2"' in view
    assert "A43" not in view


def test_excel_freezes_from_the_top(workspace, charted, ui, xl, monkeypatch):
    monkeypatch.setattr(state, "auto_mode", True)
    excel_look.tool_format_excel("schools.xlsx", "A1:B1", freeze="A2")
    window = excel_xl._app.api.ActiveWindow
    assert window.ScrollRow == 1 and window.FreezePanes is True
