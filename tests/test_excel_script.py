"""run_python_excel: what a script may not do, and the run itself (backup first, the changes shown)."""

import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from coding_agent import backups, state
from coding_agent.common import ToolError
from coding_agent.tools import documents, excel_script, excel_xl
from coding_agent.tools.excel_script import check_script, tool_run_python_excel

import fake_xlwings

TESTS = Path(__file__).parent


@pytest.fixture
def scripts_on(monkeypatch, tmp_path):
    monkeypatch.setattr(backups, "BACKUP_HOME", tmp_path / "backups")
    monkeypatch.setattr(excel_script, "EXCEL_SCRIPTS", True)
    monkeypatch.setitem(sys.modules, "xlwings", fake_xlwings)
    monkeypatch.setattr(fake_xlwings, "apps", [])
    monkeypatch.setattr(excel_xl, "_app", None)
    monkeypatch.setattr(excel_xl, "_available", None)
    monkeypatch.setattr(documents, "EXCEL_BACKEND", "xlwings")


@pytest.fixture
def book(workspace):
    wb = openpyxl.Workbook()
    wb.active.title = "Data"
    for row in [["Supplier", "Amount"], ["Acme", 120], ["Beta", 80]]:
        wb.active.append(row)
    wb.save(workspace / "book.xlsx")
    return workspace / "book.xlsx"


def excel_does(change):
    """An execute() that plays Excel: applies `change` to the workbook with openpyxl, as Excel would save it."""
    def execute(p, folder, code, timeout):
        wb = openpyxl.load_workbook(p)
        output = change(wb) or ""
        wb.save(p)
        return subprocess.CompletedProcess([], 0, output, "")
    return execute


# --- the check, before anything runs

@pytest.mark.parametrize("code", [
    "ws = book.sheets['Data']\nws.range('C1').value = 'Total'\nws.range('C2').formula = '=B2*1.2'",
    "import math, datetime\nfrom collections import Counter\nprint(math.floor(2.5))",
    "ws = book.sheets[0]\nchart = ws.charts[0]\nchart.api[1].HasTitle = True\nprint(getattr(ws, 'name'))",
    "book.sheets['Data'].range('A1').value = 'Pays | Ville'",  # a | in text is not a DDE formula
])
def test_allowed_scripts(code):
    assert check_script(code) == []


@pytest.mark.parametrize("code, refused", [
    ("book.app.api.Run('Macro1')", ".app"),
    ("book.api.Application.Run('Macro1')", ".Application"),
    ("ws = book.sheets[0]\nws.api.Evaluate('=1')", ".Evaluate"),
    ("book.api.VBProject.VBComponents.Add(1)", ".VBProject"),
    ("book.macro('Macro1')()", ".macro"),
    ("book.api.RunAutoMacros(1)", ".RunAutoMacros"),
    ("book.sheets[0].api.ExecuteExcel4Macro('EXEC(\"calc\")')", ".ExecuteExcel4Macro"),
    ("book.save('C:/elsewhere.xlsx')", ".save"),
    ("book.api.SaveAs('C:/elsewhere.xlsm')", ".SaveAs"),
    ("book.sheets[0].to_pdf('out.pdf')", ".to_pdf"),
    ("book.api.Parent.Workbooks.Open('C:/other.xlsx')", ".Parent"),
    ("book.sheets[0].api.QueryTables.Add('URL;https://x', None)", ".QueryTables"),
    ("import os\nos.system('calc')", "import os"),
    ("import subprocess", "import subprocess"),
    ("import win32com.client", "import win32com.client"),
    ("from ctypes import windll", "import ctypes"),
    ("__import__('os').system('calc')", "__import__"),
    ("eval('1+1')", "eval"),
    ("exec('print(1)')", "exec"),
    ("open('C:/x.txt', 'w').write('x')", "open"),
    ("name = 'R' + 'un'\ngetattr(book.api, name)('M')", "getattr with a computed name"),
    ("getattr(book.api, 'Run')('M')", "getattr(..., 'Run')"),
    ("book.api._oleobj_.Invoke(1, 0, 1, 1)", "._oleobj_"),
    ("().__class__.__base__.__subclasses__()", ".__class__"),
    ("book.sheets[0].range('A1').formula = \"=cmd|' /c calc'!A0\"", "starts programs or calls out"),
    ("book.sheets[0].range('A1').formula = '=WEBSERVICE(\"https://x\")'", "starts programs or calls out"),
    ("def f(:\n  pass", "syntax error"),
])
def test_refused_scripts(code, refused):
    problems = check_script(code)
    assert any(refused in p for p in problems), problems


# --- the run

def test_off_unless_switched_on(workspace, book):
    with pytest.raises(ToolError, match="AGENT_EXCEL_SCRIPTS=on"):
        tool_run_python_excel("book.xlsx", "print(1)")


def test_needs_excel_and_a_permissive_application(workspace, book, scripts_on, monkeypatch):
    monkeypatch.setattr(state, "excel_allow_format", False)
    with pytest.raises(ToolError, match="not allowed in this application"):
        tool_run_python_excel("book.xlsx", "print(1)")
    monkeypatch.setattr(state, "excel_allow_format", True)
    monkeypatch.setattr(documents, "EXCEL_BACKEND", "openpyxl")
    with pytest.raises(ToolError, match="needs Excel"):
        tool_run_python_excel("book.xlsx", "print(1)")


def test_a_refused_script_never_runs(workspace, book, scripts_on, monkeypatch, ui):
    monkeypatch.setattr(excel_script, "execute", lambda *a: pytest.fail("ran"))
    with pytest.raises(ToolError, match=r"(?s)nothing was run:.*- line 1: \.Run\n- line 1: \.app"):
        tool_run_python_excel("book.xlsx", "book.app.api.Run('M')")
    assert not ui.of("confirm") and not backups.versions("book.xlsx")


def test_declined_script_changes_nothing(workspace, book, scripts_on, monkeypatch, ui):
    monkeypatch.setattr(excel_script, "execute", lambda *a: pytest.fail("ran"))
    ui.answers = ["no", "use edit_excel"]
    with pytest.raises(ToolError, match="declined.*use edit_excel"):
        tool_run_python_excel("book.xlsx", "print(1)")
    assert not backups.versions("book.xlsx")


def test_runs_on_the_workbook_after_a_backup_and_shows_the_changes(workspace, book, scripts_on, monkeypatch, ui):
    before = book.read_bytes()

    def change(wb):
        assert backups.versions("book.xlsx")  # the backup is taken before the script runs
        wb["Data"]["C1"], wb["Data"]["C2"] = "Total", "=B2*1.2"
        wb.create_sheet("Notes")
        return "done\n"

    monkeypatch.setattr(excel_script, "execute", excel_does(change))
    ui.answers = ["yes"]
    result = tool_run_python_excel("book.xlsx", "...")
    title, rows = ui.of("cells")[0][1:3]
    assert title == "The script changed book.xlsx (2 cell(s))"
    assert rows == [("Data!C1", "", "Total", None), ("Data!C2", "", "=B2*1.2", None)]
    assert ui.of("panel")[-1][2] == ["new sheet Notes"]
    assert "2 cell(s); new sheet Notes" in result and "restore_backup" in result and "done" in result
    assert openpyxl.load_workbook(book)["Data"]["C1"].value == "Total"
    assert backups.versions("book.xlsx")[0]["id"] and documents.tool_restore_backup  # undo is possible
    ui.answers = ["yes"]
    documents.tool_restore_backup("book.xlsx", backups.versions("book.xlsx")[0]["id"])
    assert book.read_bytes() == before


def test_autonomous_mode_runs_without_asking(workspace, book, scripts_on, monkeypatch, ui):
    monkeypatch.setattr(state, "auto_mode", True)
    monkeypatch.setattr(excel_script, "execute", excel_does(lambda wb: None))
    assert "changed nothing" in tool_run_python_excel("book.xlsx", "print(1)")
    assert not ui.of("confirm")


def test_a_failed_script_reports_whether_the_workbook_changed(workspace, book, scripts_on, monkeypatch, ui):
    monkeypatch.setattr(state, "auto_mode", True)
    monkeypatch.setattr(excel_script, "execute",
                        lambda *a: subprocess.CompletedProcess([], 1, "", "NameError: name 'ws' is not defined"))
    with pytest.raises(ToolError, match="(?s)exit code 1.*was not changed.*NameError"):
        tool_run_python_excel("book.xlsx", "ws.range('A1').value = 1")


def test_macros_added_are_undone(workspace, book, scripts_on, monkeypatch, ui):
    monkeypatch.setattr(state, "auto_mode", True)
    before = book.read_bytes()
    real_parts = excel_script.parts
    # after the run, the workbook holds macros it did not hold before
    monkeypatch.setattr(excel_script, "parts", lambda p: {**real_parts(p), "macros": 1 if p == book else 0})
    monkeypatch.setattr(excel_script, "execute", excel_does(lambda wb: None))
    with pytest.raises(ToolError, match="added macros"):
        tool_run_python_excel("book.xlsx", "print(1)")
    assert book.read_bytes() == before


def test_the_real_runner_runs_in_the_sandbox(workspace, book, scripts_on, monkeypatch, ui, tmp_path):
    """The runner and the guard, for real, with fake_xlwings as Excel. Excel itself saves from outside
    Python; the fake saves from Python, so the sandbox blocks its save in the workspace -- which is
    what it must do to any Python code there."""
    shim = tmp_path / "shim"
    shim.mkdir()
    (shim / "xlwings.py").write_text(f"import sys\nsys.path.insert(0, {str(TESTS)!r})\nfrom fake_xlwings import *\n")
    monkeypatch.setenv("PYTHONPATH", str(shim))
    monkeypatch.setattr(state, "auto_mode", True)
    before = book.read_bytes()
    with pytest.raises(ToolError) as refused:
        tool_run_python_excel("book.xlsx", "ws = book.sheets['Data']\nprint('rows:', ws.range('B2').value)\n"
                                           "ws.range('C1').value = 'Total'")
    message = str(refused.value)
    assert "rows: 120" in message  # the script ran, with `book`
    assert "Blocked by the coding agent" in message and "was not changed" in message
    assert book.read_bytes() == before
